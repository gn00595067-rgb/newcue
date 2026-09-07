# -*- coding: utf-8 -*-
"""
樣式母版轉印機制（§2.3 / 附錄 A）

把 assets/style_masters/*.xlsx（只有樣式、無值）的欄寬/列高/合併/列印設定/逐格樣式
整份轉印到新分頁（日期欄依實際天數展開），代理商版面即與 0902 範本 100% 相同；
渲染器只負責填值與動態覆寫（六日底色/月份/日期/回饋文字/Logo）。
"""
import os
from copy import copy
from functools import lru_cache

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

MASTER_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "assets", "style_masters")

# key: (FIRST 第一日期欄, LAST 最後日期欄, MAXROW, DATA_LAST_ROW 資料最後列, PRINT_LAST_ROW)
PARAMS = {
    "ag_2008_fam":  (9, 29, 24, 14, 24),
    "ag_2008_wjf":  (9, 38, 24, 15, 24),
    "ag_carat_fam": (11, 32, 32, 11, 32),
    "ag_carat_wjf": (11, 21, 32, 9, 32),
}


@lru_cache(maxsize=None)
def _master_ws(key):
    return load_workbook(os.path.join(MASTER_DIR, f"{key}.xlsx")).worksheets[0]


def _cp(src, dst):
    dst.font = copy(src.font)
    dst.border = copy(src.border)
    dst.fill = copy(src.fill)
    dst.alignment = copy(src.alignment)
    dst.number_format = src.number_format


class StyleMaster:
    def __init__(self, key):
        self.key = key
        self.first, self.last, self.maxrow, self.data_last, self.print_last = PARAMS[key]
        self.tw = _master_ws(key)

    def apply(self, ws, n_days):
        """把母版樣式轉印到 ws（日期欄展開成 n_days）；回傳 {'first','last'}。"""
        tw = self.tw
        FIRST, LAST, MAXROW = self.first, self.last, self.maxrow
        out_last = FIRST + n_days - 1

        # 1) 欄寬（固定欄；column_dimensions 一筆可涵蓋 min..max，需展開）
        for d in tw.column_dimensions.values():
            if d.width is None:
                continue
            for c in range(d.min, d.max + 1):
                if c < FIRST:
                    ws.column_dimensions[get_column_letter(c)].width = d.width
        day_w = tw.column_dimensions[get_column_letter(FIRST)].width
        if day_w:
            for j in range(n_days):
                ws.column_dimensions[get_column_letter(FIRST + j)].width = day_w
        if tw.sheet_format.defaultRowHeight:
            ws.sheet_format.defaultRowHeight = tw.sheet_format.defaultRowHeight
        if tw.sheet_format.defaultColWidth:
            ws.sheet_format.defaultColWidth = tw.sheet_format.defaultColWidth

        # 2) 列高
        for r in range(1, MAXROW + 1):
            h = tw.row_dimensions[r].height
            if h:
                ws.row_dimensions[r].height = h

        # 3) 逐格樣式：固定欄照抄；日期欄 j → 第一天/中間/最後一天三種樣式
        for r in range(1, MAXROW + 1):
            for c in range(1, FIRST):
                _cp(tw.cell(row=r, column=c), ws.cell(row=r, column=c))
            for j in range(n_days):
                src_c = FIRST if j == 0 else (LAST if j == n_days - 1 else FIRST + 1)
                _cp(tw.cell(row=r, column=src_c), ws.cell(row=r, column=FIRST + j))

        # 4) 合併
        for rng in tw.merged_cells.ranges:
            if rng.min_col < FIRST:
                mc = rng.max_col if rng.max_col < FIRST else out_last
                ws.merge_cells(start_row=rng.min_row, start_column=rng.min_col,
                               end_row=rng.max_row, end_column=mc)
            elif rng.min_row <= self.data_last:      # 資料列的日期區合併 → 映射 FIRST..out_last
                ws.merge_cells(start_row=rng.min_row, start_column=FIRST,
                               end_row=rng.max_row, end_column=out_last)
            # else：日期區內範本殘留合併（如 2008 I15:Z15）→ 略過

        # 5) 列印
        ws.page_setup.orientation = tw.page_setup.orientation
        ws.page_setup.paperSize = tw.page_setup.paperSize
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1
        ws.page_margins = copy(tw.page_margins)
        ws.print_options.horizontalCentered = tw.print_options.horizontalCentered
        ws.print_options.verticalCentered = tw.print_options.verticalCentered
        ws.print_title_rows = tw.print_title_rows
        ws.sheet_view.showGridLines = False
        ws.sheet_view.zoomScale = tw.sheet_view.zoomScale
        ws.print_area = f"A1:{get_column_letter(out_last)}{self.print_last}"
        return {"first": FIRST, "last": out_last}
