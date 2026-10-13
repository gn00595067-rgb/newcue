# -*- coding: utf-8 -*-
"""一鍵CUE 選聲活／鉑霖：轉一般 CUE rows 後用一般 CUE renderer 輸出（docs/specs/格式選擇擴充.md）。"""
import io
from datetime import date, timedelta

import pytest
from openpyxl import load_workbook

import simple_config as sc
import simple_general as sg
import simple_model as sm
from fixtures_simple import sheetdata

PERIODS = [(date(2026, 9, 21), 7), (date(2026, 9, 21), 35)]   # 單月／跨月超過 31 天


def _model(key, budget=250000, start=date(2026, 9, 21), days=14):
    end = start + timedelta(days=days - 1)
    return sm.build_model(key, budget, start, end, client="客戶甲", tax_id="12345678",
                          product="產品乙", data=sheetdata())


@pytest.mark.parametrize("key", sc.SUBSIDIARY_COMBOS)
def test_rows_match_model(key):
    m = _model(key)
    for sheet in m.sheets:
        rows = sg.sheet_to_rows(sheet)
        # 每平台 Package-cost 合計 = 預算
        pkgs = {r["media"]: r["nat_pkg_display"] for r in rows}
        assert sum(pkgs.values()) == sheet.budget
        # 檔次與簡易模式一致
        assert sum(r["spots"] for r in rows) == sum(r.spots for b in sheet.blocks for r in b.rows)
        for r in rows:
            assert sum(r["schedule"]) == r["spots"]
            if r["media"] == "家樂福":
                assert r["region"] in ("全省量販", "全省超市")


@pytest.mark.parametrize("fmt", ["聲活", "鉑霖"])
@pytest.mark.parametrize("key", sc.SUBSIDIARY_COMBOS)
@pytest.mark.parametrize("start,days", PERIODS)
def test_render_xlsx(fmt, key, start, days):
    m = _model(key, start=start, days=days)
    wb = load_workbook(io.BytesIO(sg.render_xlsx(m, fmt, sales_person="王小明")))
    names = wb.sheetnames
    assert len(names) == len(set(names))
    for sheet in m.sheets:
        assert any(n.startswith(f"{sheet.seconds}秒") for n in names)
    if days <= 31:
        assert len(names) == len(m.sheets)
    title = "聲活數位-媒體計劃排程表" if fmt == "聲活" else "鉑霖行動行銷-媒體計劃排程表"
    for ws in wb.worksheets:
        assert title in str(ws["A1"].value or "")


@pytest.mark.parametrize("fmt", ["聲活", "鉑霖"])
def test_render_htmls(fmt):
    m = _model("sub_qp_wjf")
    htmls = sg.render_htmls(m, fmt)
    assert len(htmls) == len(m.sheets)
    assert all("客戶甲" in h for _, h in htmls)


def test_existing_renderer_unchanged_without_wb():
    """未傳 wb 時 generate_excel_from_scratch 仍回傳 bytes、單一工作表 Schedule。"""
    from excel_renderer import generate_excel_from_scratch
    m = _model("sub_qp_fv")
    rows = sg.sheet_to_rows(m.sheets[0])
    b = generate_excel_from_scratch("東吳", m.start, m.end, "客", "", "品", rows, ["1.x"],
                                    m.sheets[0].budget, 0, "", None)
    wb = load_workbook(io.BytesIO(b))
    assert wb.sheetnames == ["Schedule"]
