# -*- coding: utf-8 -*-
"""健康視CUE（宜歆 10/29 #6；spec docs/specs/健康視CUE.md）：曝光／定價計算、備註、Excel 值版。"""
import io
from datetime import date

from openpyxl import load_workbook

import health_cue as h
import utils

D = dict(format="鉑霖", client="小檸檬媒體整合行銷有限公司", advertiser="", tax_id="55683289", ad_name="",
         sales="蔡宜歆", seconds=15, start=date(2026, 11, 23), end=date(2026, 12, 22), days="週一至週五",
         budget=300000, sign_deadline=date(2026, 10, 30), billing_month="116/1月", payment_date=date(2027, 2, 22))


def test_calc_matches_reference():
    assert h.calc(15, 300000) == {"impressions": 460800, "list": 450000, "net": 300000, "vat": 15000, "gross": 315000}
    assert h.calc(10, 50000)["impressions"] == 115200          # 宜歆範例：384000/166667*50000
    for s, (_l, _d, proj_tax) in h.DEFAULT_PRICING.items():   # 一個專案價＝38.4 萬次
        assert h.calc(s, proj_tax / 1.05)["impressions"] == 384000


def test_calc_uses_given_pricing():
    p = {15: (393750, 300000, 525000)}
    assert h.calc(15, 250000, p)["impressions"] == 192000


def test_remarks_and_defaults():
    rem = h.remarks(date(2026, 10, 30), "116/1月", date(2027, 2, 22))
    assert rem[0] == ("1. 請於 2026/10/30 (五) 下午1500前回簽。", "FF0000")
    assert rem[4][0] == "5. 付款兌現日期：116/2/22前兌現。" and rem[5][1] == "0000FF" and len(rem) == 6
    assert h.default_billing(date(2026, 12, 22)) == "116/1月"
    assert h.default_payment(date(2026, 12, 22)) == date(2027, 1, 31)
    assert h.default_payment(date(2026, 11, 5)) == date(2026, 12, 31)


def test_excel_values():
    ws = load_workbook(io.BytesIO(h.build_excel(D))).active
    assert ws["B1"].value.startswith("鉑霖行動行銷") and ws["C4"].value == "鉑霖行動行銷 - 蔡宜歆"
    assert ws["G9"].value == "11/23-12/22\n執行週一至週五" and ws["F9"].value.startswith("15秒")
    assert (ws["H9"].value, ws["I9"].value, ws["J9"].value) == (460800, 450000, 300000)   # 值版，無公式
    assert (ws["J12"].value, ws["J13"].value) == (15000, 315000)
    assert ws["C11"].value == D["client"] and ws["B20"].value.startswith("6. 若實際曝光次數")
    ws2 = load_workbook(io.BytesIO(h.build_excel(dict(D, format="聲活", advertiser="廣告主乙", days="每日")))).active
    assert ws2["B1"].value.startswith("聲活數位") and ws2["C11"].value == "廣告主乙" and ws2["G9"].value.endswith("執行每日")


def test_platform_text_for_ragic():
    rows = [{"media": "全民健康視", "region": "全台"}]
    assert utils.build_platform_text(rows) == "全民健康視"
    assert utils.build_platform_detail_text(rows, {}) == "全民健康視"


def test_filename():
    assert h.filename("客戶甲", 15, 300000, "宜", today=date(2026, 10, 30)) == "1030 客戶甲 全民健康視(15秒) 30萬-宜.xlsx"
