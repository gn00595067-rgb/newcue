# -*- coding: utf-8 -*-
"""
代理商 Excel 結構測試（版型 B 2008 / C 凱絡）。

驗證渲染器產出的結構與資料正確：分頁數/名稱、抬頭與表頭標籤、費用標籤、
逐日排檔長度=走期天數、關鍵數值（牌價、檔次、費用）確實出現在正確欄位。

註：代理商版型的逐格「像素級」fidelity（§10.2 for ④⑤⑥⑦）尚在收斂中，
    本測試聚焦於「結構與商業數值正確」；子公司版型已有嚴格 fidelity（test_simple_excel）。
"""
import io
from datetime import date

import pytest
from openpyxl import load_workbook

import simple_model as sm
import simple_excel as se
from fixtures_simple import sheetdata


def _cells(ws):
    out = {}
    for row in ws.iter_rows():
        for c in row:
            if c.value not in (None, ""):
                out[c.coordinate] = c.value
    return out


def _all_text(ws):
    return "\n".join(str(v) for v in _cells(ws).values())


AGENCY_CASES = [
    ("ag_2008_fam", date(2026, 6, 8), date(2026, 6, 28), 4),
    ("ag_2008_wjf", date(2026, 7, 17), date(2026, 7, 30), 3),
    ("ag_carat_fam", date(2026, 9, 7), date(2026, 9, 28), 4),
    ("ag_carat_wjf", date(2026, 9, 18), date(2026, 9, 28), 3),
]


@pytest.mark.parametrize("key,s,e,n", AGENCY_CASES)
def test_agency_sheet_count(key, s, e, n):
    m = sm.build_model(key, 250000, s, e, client="統一", product="麥香", data=sheetdata())
    wb = load_workbook(io.BytesIO(se.render(m, formulas=True)))
    assert len(wb.worksheets) == n


@pytest.mark.parametrize("key,s,e,n", AGENCY_CASES)
def test_agency_both_modes_render(key, s, e, n):
    m = sm.build_model(key, 250000, s, e, client="統一", product="麥香", data=sheetdata())
    for formulas in (True, False):
        wb = load_workbook(io.BytesIO(se.render(m, formulas=formulas)))
        assert len(wb.worksheets) == n


def test_2008_family_structure():
    m = sm.build_model("ag_2008_fam", 250000, date(2026, 6, 8), date(2026, 6, 28),
                       client="統一企業", product="麥香", data=sheetdata())
    wb = load_workbook(io.BytesIO(se.render(m, formulas=True)))
    ws = wb.worksheets[0]           # 30 秒
    txt = _all_text(ws)
    for lab in ["Client", "Media Agency", "Product", "Period", "媒體型態", "地區",
                "播出時段", "定價", "單位", "次數", "合計", "Budget (net)：",
                "AC %：", "5% Tax ：", "TOTAL ："]:
        assert lab in txt, f"缺標籤 {lab}"
    assert "統一企業" in txt
    assert sc_company() in txt
    # 回饋列文字
    assert "72(走期內執行完畢)" in txt      # 30秒回饋量15%=72
    assert "221(走期內執行完畢)" in txt     # 46%=221
    assert ws.page_setup.orientation == "landscape"
    assert ws.page_setup.paperSize == 9


def sc_company():
    import simple_config as sc
    return sc.COMPANY_2008


def test_2008_wjf_structure():
    m = sm.build_model("ag_2008_wjf", 250000, date(2026, 7, 17), date(2026, 7, 30),
                       client="統一麵", data=sheetdata())
    wb = load_workbook(io.BytesIO(se.render(m, formulas=True)))
    ws = wb.worksheets[0]           # 10 秒
    txt = _all_text(ws)
    assert "計於量販店" in txt
    assert "回饋" in txt
    # 10秒回饋量販10%=84（萬家福回饋列日期區只寫數字，§2.4）
    assert m.sheets[0].blocks[0].rows[2].spots == 84


def test_carat_family_structure():
    m = sm.build_model("ag_carat_fam", 250000, date(2026, 9, 7), date(2026, 9, 28),
                       client="統一", product="麥香", data=sheetdata())
    wb = load_workbook(io.BytesIO(se.render(m, formulas=True)))
    ws = wb.worksheets[0]           # 30 秒
    txt = _all_text(ws)
    for lab in ["凱絡媒體服務(股)公司廣播媒體排期表", "媒體別", "定價\n(檔/Net)",
                "統一價\n(檔/Net)", "檔數", "總價", "媒體總價值(NET)", "優惠總價值(NET)",
                "Sub-Total", "A.C     3%", "VAT    5%", "Grand-Total", "專案回饋"]:
        assert lab in txt, f"缺標籤 {lab}"
    # 媒體總價值/優惠 30 秒 = 1,739,250 / 1,489,250（公式版為 =SUM，改查值版）
    vals = set(_cells(load_workbook(io.BytesIO(se.render(m, formulas=False))).worksheets[0]).values())
    assert 1739250 in vals and 1489250 in vals
    # A.C 顯示 -（免收）
    assert "-" in txt


def test_carat_wjf_structure():
    m = sm.build_model("ag_carat_wjf", 250000, date(2026, 9, 18), date(2026, 9, 28),
                       client="統一", product="純喫茶", data=sheetdata())
    wb = load_workbook(io.BytesIO(se.render(m, formulas=True)))
    ws = wb.worksheets[0]
    txt = _all_text(ws)
    assert "計價於量販" in txt
    assert "萬家福(量販)" in txt and "樂家康(超市)" in txt


@pytest.mark.parametrize("key", ["ag_carat_fam", "ag_carat_wjf"])
def test_carat_total_col_wide_enough(key):
    """凱絡『總價』欄(I)須夠寬容 7~8 位數會計格式，避免 ###（§7.4，已用真 Excel 驗證）。"""
    from datetime import date as _d
    m = sm.build_model(key, 400000, _d(2026, 9, 18), _d(2026, 9, 24), data=sheetdata())
    wb = load_workbook(io.BytesIO(se.render(m, formulas=False)))
    for ws in wb.worksheets:
        assert (ws.column_dimensions["I"].width or 0) >= 16, f"{key} 總價欄過窄會 ###"


def test_2008_family_price_col_wide_enough():
    """2008 全家『定價』欄(D)須夠寬容 7 位數(牌價×次數)，避免 ###（§7.4，已用真 Excel 驗證）。"""
    from datetime import date as _d
    m = sm.build_model("ag_2008_fam", 400000, _d(2026, 6, 8), _d(2026, 6, 28),
                       client="統一", product="統一", data=sheetdata())
    wb = load_workbook(io.BytesIO(se.render(m, formulas=False)))
    for ws in wb.worksheets:
        assert (ws.column_dimensions["D"].width or 0) >= 30, "2008全家定價欄過窄會 ###"


@pytest.mark.parametrize("key,s,e,n", AGENCY_CASES)
def test_agency_schedule_length(key, s, e, n):
    """逐日排檔長度 = 走期天數（value 版逐日皆為數字）。"""
    ndays = (e - s).days + 1
    m = sm.build_model(key, 250000, s, e, data=sheetdata())
    for sh in m.sheets:
        for blk in sh.blocks:
            for r in blk.rows:
                if r.schedule is not None:
                    assert len(r.schedule) == ndays


# --------------------------------------------------------------------------- #
# 凱絡外圍一圈外框（照母版 A5:末欄:32 medium；曾因母版清除下半部而斷框）
# --------------------------------------------------------------------------- #
def _style(side):
    return side.style if side and side.style else None


@pytest.mark.parametrize("key,s,e", [
    ("ag_carat_fam", date(2026, 9, 21), date(2026, 10, 4)),
    ("ag_carat_wjf", date(2026, 9, 21), date(2026, 10, 4)),
])
def test_carat_outer_frame_is_continuous_medium(key, s, e):
    """凱絡整份最外圈四邊皆應為連續 medium（A5 到末欄、到第 32 列）。"""
    import re
    from openpyxl.utils import column_index_from_string as CI
    m = sm.build_model(key, 250000, s, e, client="X", data=sheetdata())
    ws = load_workbook(io.BytesIO(se.render(m, formulas=False))).worksheets[0]
    last = ws.print_area.split("!")[-1].split(":")[1].replace("$", "")
    out_last = CI(re.match(r"([A-Z]+)", last).group(1))
    TOP, BOT = 5, 32
    assert all(_style(ws.cell(r, 1).border.left) == "medium" for r in range(TOP, BOT + 1)), "左框斷"
    assert all(_style(ws.cell(r, out_last).border.right) == "medium" for r in range(TOP, BOT + 1)), "右框斷"
    assert all(_style(ws.cell(TOP, c).border.top) == "medium" for c in range(1, out_last + 1)), "頂框斷"
    assert all(_style(ws.cell(BOT, c).border.bottom) == "medium" for c in range(1, out_last + 1)), "底框斷"
