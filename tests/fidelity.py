# -*- coding: utf-8 -*-
"""
範本 fidelity 比對工具（§10.2）—— 供 test_simple_excel 使用。

比對「我方輸出 sheet」vs「cueexample 範本 sheet」，回傳差異清單。
等價正規化（視覺無差異者視為相同）：
  - 字色：theme1 / 純黑 / 無 → 一律「預設黑」
  - 底色：theme0 / 純白 / 無 → 一律「無底色」（只保留有意義填色，如六日 FFFF99）
  - 數字格式：去除引號逸出差異（"店" == ""店""）；文字型儲存格不比格式
  - 文字值：忽略空白差異（範本常有多餘換行/尾空白）
數字/公式儲存格只比「有無值 + 格式」，不比數值（範本檔次為人工挑選）。
內部欄（檔次欄 V 右側）不比對。
"""
import re
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, column_index_from_string


def _fontcolor(c):
    if c is None:
        return "BLACK"
    rgb = getattr(c, "rgb", None)
    if isinstance(rgb, str) and rgb not in ("00000000", "FF000000"):
        return rgb
    t = getattr(c, "theme", None)
    if t == 1:
        return "BLACK"
    return f"theme{t}" if t is not None else "BLACK"


def _fillc(cell):
    f = cell.fill
    if f is None or f.patternType != "solid":
        return "NONE"
    fg = f.fgColor
    rgb = getattr(fg, "rgb", None)
    if isinstance(rgb, str) and rgb not in ("00000000", "FFFFFFFF"):
        return rgb
    t = getattr(fg, "theme", None)
    if isinstance(t, int) and t != 0:      # theme0=白≈NONE；非 int(auto/特殊) 一律 NONE
        return f"theme{t}"
    return "NONE"


def _bd(cell):
    b = cell.border
    def s(x):
        st = x.style if (x and x.style) else "."
        return "thin" if st == "hair" else st   # 內部格線 hair 一律視同 thin（我方升級為 thin 讓格線連續）
    return f"{s(b.top)},{s(b.bottom)},{s(b.left)},{s(b.right)}"


# --- 框線「有效邊」比對（§7）：hair≡thin、seal 造成的鄰格重複不算差異 --------------- #
def _border_level(style):
    """框線粗細分級：0=無、1=細(hair/thin/虛線)、2=粗(medium/thick)、3=double。"""
    if not style or style == ".":
        return 0
    if style in ("hair", "thin", "dotted", "dashed", "dashDot", "dashDotDot"):
        return 1
    if style == "double":
        return 3
    return 2   # medium / thick / mediumDashed ...


def _side_style(cell, side):
    return getattr(getattr(cell.border, side), "style", None)


def _merge_map(ws):
    """(r,c) → 所屬合併範圍 key；非合併格不列入。"""
    mp = {}
    for rng in ws.merged_cells.ranges:
        k = (rng.min_row, rng.min_col, rng.max_row, rng.max_col)
        for rr in range(rng.min_row, rng.max_row + 1):
            for cc in range(rng.min_col, rng.max_col + 1):
                mp[(rr, cc)] = k
    return mp


def _merge_internal(mmap, r, c, side):
    """該格某邊是否為『合併格內部邊』（與鄰格屬同一合併範圍）→ Excel 不顯示。"""
    nb = {"top": (r - 1, c), "bottom": (r + 1, c),
          "left": (r, c - 1), "right": (r, c + 1)}[side]
    k = mmap.get((r, c))
    return k is not None and mmap.get(nb) == k


def _eff_level(ws, r, c, side, maxrow):
    """該格某邊的『有效』框線分級 = max(自身該邊, 相鄰格的貼合邊)。
    Excel 兩格之間的線只要任一側宣告即會顯示；我方 seal 會在兩側都畫，範本常只畫一側。"""
    own = _side_style(ws.cell(row=r, column=c), side)
    nb = None
    if side == "top" and r > 1:
        nb = _side_style(ws.cell(row=r - 1, column=c), "bottom")
    elif side == "bottom" and r < maxrow:
        nb = _side_style(ws.cell(row=r + 1, column=c), "top")
    elif side == "left" and c > 1:
        nb = _side_style(ws.cell(row=r, column=c - 1), "right")
    elif side == "right":
        nb = _side_style(ws.cell(row=r, column=c + 1), "left")
    return max(_border_level(own), _border_level(nb))


def _norm_nf(s):
    return (s or "General").replace('"', "")


def _norm_ws(v):
    return re.sub(r"\s", "", str(v))


def _attrs(cell):
    a = cell.alignment
    return dict(sz=cell.font.size, bold=bool(cell.font.bold),
                fc=_fontcolor(cell.font.color), bg=_fillc(cell),
                nf=_norm_nf(cell.number_format),
                al=f"{a.horizontal}/{a.vertical}/{'W' if a.wrap_text else '.'}",
                bd=_bd(cell))


def _is_text(v):
    return isinstance(v, str) and not v.startswith("=")


def _find_cv(tw):
    """檔次欄索引（找表頭『檔次』）；找不到用最大欄。"""
    for c in range(1, tw.max_column + 1):
        if tw.cell(row=7, column=c).value == "檔次":
            return c
    return tw.max_column


def compare(tmpl_ws, out_ws, cv=None):
    diffs = []
    tm = {str(r) for r in tmpl_ws.merged_cells.ranges}
    om = {str(r) for r in out_ws.merged_cells.ranges}
    for m in sorted(tm - om):
        diffs.append(f"MERGED-MISSING {m}")
    for m in sorted(om - tm):
        diffs.append(f"MERGED-EXTRA {m}")
    def _range_only(pa):
        # 去掉工作表名前綴（範本 10 秒分頁名尾帶空白），只比列印範圍
        return str(pa).split("!")[-1] if pa else pa
    for attr, tv, ov in [
        ("orientation", tmpl_ws.page_setup.orientation, out_ws.page_setup.orientation),
        ("paper", tmpl_ws.page_setup.paperSize, out_ws.page_setup.paperSize),
        ("print_area", _range_only(tmpl_ws.print_area), _range_only(out_ws.print_area)),
    ]:
        if str(tv) != str(ov):
            diffs.append(f"PRINT-{attr} tmpl={tv} ours={ov}")
    cv = cv or _find_cv(tmpl_ws)
    for letter, dim in tmpl_ws.column_dimensions.items():
        if dim.width and column_index_from_string(letter) <= cv:
            ov = out_ws.column_dimensions[letter].width
            if ov is None or abs(ov - dim.width) > 0.5:
                diffs.append(f"WIDTH {letter} tmpl={dim.width:.1f} ours={ov}")
    for idx, dim in tmpl_ws.row_dimensions.items():
        if dim.height:
            ov = out_ws.row_dimensions[idx].height
            if ov is None or abs(ov - dim.height) > 0.5:
                diffs.append(f"HEIGHT {idx} tmpl={dim.height:.1f} ours={ov}")
    for r in range(1, tmpl_ws.max_row + 1):
        for c in range(1, cv + 1):
            tc = tmpl_ws.cell(row=r, column=c)
            if tc.value is None or tc.value == "":
                continue
            oc = out_ws.cell(row=r, column=c)
            loc = f"{get_column_letter(c)}{r}"
            ta, oa = _attrs(tc), _attrs(oc)
            txt = _is_text(tc.value)
            if txt:
                if _norm_ws(tc.value) != _norm_ws(oc.value):
                    diffs.append(f"VALUE {loc}: tmpl={tc.value!r} ours={oc.value!r}")
            elif oc.value is None or oc.value == "":
                diffs.append(f"EMPTY {loc}: tmpl has {tc.value!r}")
            keys = ("sz", "bold", "fc", "bg", "al", "bd") if txt else \
                   ("sz", "bold", "fc", "bg", "nf", "al", "bd")
            for k in keys:
                if ta[k] != oa[k]:
                    diffs.append(f"{k.upper()} {loc}: tmpl={ta[k]} ours={oa[k]}")
    return diffs


def load(path):
    return load_workbook(path)


# =========================================================================== #
# strict 比對（§5 / 附錄 B）：空格也比 border/fill；日期欄映射；列印設定
# =========================================================================== #
def _print_last_row(ws):
    pa = ws.print_area
    if not pa:
        return ws.max_row
    m = re.search(r"\$?[A-Z]+\$?(\d+)\s*$", str(pa).split("!")[-1])
    return int(m.group(1)) if m else ws.max_row


def compare_strict(tw, ow, *, first, out_first, tmpl_days, out_days, right_cols=0,
                   print_last=None, data_last=None, identity_merge=False,
                   check_borders=True):
    """
    first/out_first：第一個日期欄（1-based）；tmpl_days/out_days：範本/我方日期欄數。
    right_cols：日期欄右側固定欄數（子公司檔次欄 V = 1；代理商 0）。
    """
    diffs = []
    tmpl_last = first + tmpl_days - 1
    out_last = out_first + out_days - 1
    plast = print_last or _print_last_row(tw)

    # --- 列印設定 ---
    ps_t, ps_o = tw.page_setup, ow.page_setup
    if str(ps_t.orientation) != str(ps_o.orientation):
        diffs.append(f"PRINT-orientation tmpl={ps_t.orientation} ours={ps_o.orientation}")
    if str(ps_t.paperSize) != str(ps_o.paperSize):
        diffs.append(f"PRINT-paper tmpl={ps_t.paperSize} ours={ps_o.paperSize}")
    if str(ow.page_setup.fitToWidth) not in ("1", "None") or str(ow.page_setup.fitToHeight) not in ("1", "None"):
        # 我方一律 fitToWidth=1 fitToHeight=1
        if ow.page_setup.fitToWidth != 1 or ow.page_setup.fitToHeight != 1:
            diffs.append(f"PRINT-fit ours=W{ow.page_setup.fitToWidth}/H{ow.page_setup.fitToHeight}")
    if str(tw.print_title_rows) != str(ow.print_title_rows):
        diffs.append(f"PRINT-title_rows tmpl={tw.print_title_rows} ours={ow.print_title_rows}")
    for a in ("horizontalCentered", "verticalCentered"):
        if bool(getattr(tw.print_options, a)) != bool(getattr(ow.print_options, a)):
            diffs.append(f"PRINT-{a} tmpl={getattr(tw.print_options,a)} ours={getattr(ow.print_options,a)}")
    for a in ("left", "right", "top", "bottom"):
        tv, ov = getattr(tw.page_margins, a), getattr(ow.page_margins, a)
        if tv is not None and ov is not None and abs(tv - ov) > 0.05:
            diffs.append(f"MARGIN-{a} tmpl={tv:.2f} ours={ov:.2f}")
    if bool(tw.sheet_view.showGridLines) != bool(ow.sheet_view.showGridLines):
        diffs.append(f"gridlines tmpl={tw.sheet_view.showGridLines} ours={ow.sheet_view.showGridLines}")
    if _print_last_row(ow) != plast:
        diffs.append(f"PRINT-last-row tmpl={plast} ours={_print_last_row(ow)}")

    # --- 欄寬（固定欄，展開 min..max）---
    tw_w = {}
    for d in tw.column_dimensions.values():
        if d.width:
            for c in range(d.min, d.max + 1):
                tw_w[c] = d.width
    ow_w = {}
    for d in ow.column_dimensions.values():
        if d.width:
            for c in range(d.min, d.max + 1):
                ow_w[c] = d.width
    for c in range(1, first):
        if c in tw_w:
            ov = ow_w.get(c)
            if ov is None or abs(ov - tw_w[c]) > 0.5:
                diffs.append(f"WIDTH {get_column_letter(c)} tmpl={tw_w[c]:.1f} ours={ov}")

    # --- 列高 1..plast ---
    for r in range(1, plast + 1):
        th = tw.row_dimensions[r].height
        if th:
            oh = ow.row_dimensions[r].height
            if oh is None or abs(oh - th) > 0.5:
                diffs.append(f"HEIGHT {r} tmpl={th:.1f} ours={oh}")

    # --- 欄配對 ---
    pairs = [(c, c) for c in range(1, first)]
    if tmpl_days == out_days:
        pairs += [(first + k, out_first + k) for k in range(out_days)]
    else:
        pairs += [(first, out_first), (first + 1, out_first + 1), (tmpl_last, out_last)]
    for k in range(1, right_cols + 1):
        pairs.append((tmpl_last + k, out_last + k))

    tmap, omap = _merge_map(tw), _merge_map(ow)
    for r in range(1, plast + 1):
        for tc_col, oc_col in pairs:
            tc = tw.cell(row=r, column=tc_col)
            oc = ow.cell(row=r, column=oc_col)
            loc = f"{get_column_letter(tc_col)}{r}"
            ta, oa = _attrs(tc), _attrs(oc)
            # 框線比「有效邊」（§7）：hair≡thin、seal 造成的鄰格重複不算差異；
            # 只回報「範本有線我方沒有」與「粗細等級不同」，不回報「我方多畫」。
            # 天數不同時（範本畫整月、我方只畫走期）欄位配對不可靠，略過框線比對。
            if check_borders and out_days == tmpl_days:
                for side in ("top", "bottom", "left", "right"):
                    # 合併格內部邊在 Excel 不顯示 → 略過（範本/我方任一為內部邊皆略過）
                    if _merge_internal(tmap, r, tc_col, side) or \
                       _merge_internal(omap, r, oc_col, side):
                        continue
                    tl = _eff_level(tw, r, tc_col, side, plast)
                    ol = _eff_level(ow, r, oc_col, side, plast)
                    if tl > 0 and ol == 0:
                        diffs.append(f"BD {loc}: 缺{side}線 tmpl_lvl={tl}")
                    elif tl > 0 and ol > 0 and tl != ol:
                        diffs.append(f"BD {loc}: {side}粗細 tmpl_lvl={tl} ours_lvl={ol}")
            if ta["bg"] != oa["bg"]:
                diffs.append(f"BG {loc}: tmpl={ta['bg']} ours={oa['bg']}")
            if tc.value in (None, "") or (isinstance(tc.value, str) and _norm_ws(tc.value) == ""):
                continue    # 範本空白/純空白字元格：只比 border/fill（上面已比）
            txt = _is_text(tc.value)
            if txt and _norm_ws(tc.value) != _norm_ws(oc.value):
                diffs.append(f"VALUE {loc}: tmpl={tc.value!r} ours={oc.value!r}")
            elif not txt and (oc.value is None or oc.value == ""):
                diffs.append(f"EMPTY {loc}: tmpl has {tc.value!r}")
            keys = ("sz", "bold", "fc", "al") if txt else ("sz", "bold", "fc", "nf", "al")
            for kk in keys:
                if ta[kk] != oa[kk]:
                    diffs.append(f"{kk.upper()} {loc}: tmpl={ta[kk]} ours={oa[kk]}")

    # --- 合併 ---
    dlast = data_last if data_last is not None else plast
    off = out_last - tmpl_last

    def _map(rng):
        if rng.max_col < first:                      # 全在左側固定欄
            return (rng.min_row, rng.min_col, rng.max_row, rng.max_col)
        if rng.min_col > tmpl_last:                  # 全在右側固定欄（子公司 V）
            return (rng.min_row, rng.min_col + off, rng.max_row, rng.max_col + off)
        if rng.min_col < first:                      # 由左側跨入日期/右側（如標題 A1:V1）
            if rng.max_col > tmpl_last:
                mc = rng.max_col + off
            elif rng.max_col >= first:
                mc = out_last
            else:
                mc = rng.max_col
            return (rng.min_row, rng.min_col, rng.max_row, mc)
        if rng.min_row <= dlast:                     # 日期區內、資料列 → 映射整段
            return (rng.min_row, out_first, rng.max_row, out_last)
        return None

    if identity_merge:                               # 子公司：天數相符且無殘留合併 → 原座標
        tmpl_m = {(r.min_row, r.min_col, r.max_row, r.max_col) for r in tw.merged_cells.ranges}
    else:                                            # 代理商：用映射並略過日期區殘留合併
        tmpl_m = {_map(r) for r in tw.merged_cells.ranges if _map(r)}
    ours_m = {(r.min_row, r.min_col, r.max_row, r.max_col) for r in ow.merged_cells.ranges}
    for m in sorted(tmpl_m - ours_m):
        diffs.append(f"MERGE-MISSING {m}")
    for m in sorted(ours_m - tmpl_m):
        diffs.append(f"MERGE-EXTRA {m}")
    return diffs
