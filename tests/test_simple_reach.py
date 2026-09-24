# -*- coding: utf-8 -*-
"""預估曝光／人流測試（§6.1/§6.3）：對範本『各平台人流計算方式.xlsx』12 個數字精確相等。"""
from datetime import date

import pytest

import simple_config as sc
import simple_model as sm
import simple_reach as sr
from agency_cue import rhu
from fixtures_simple import sheetdata


# --------------------------------------------------------------------------- #
# §1.5 驗證向量：直接建一組 blocks 餵 compute_reach（不經預算分配）
# --------------------------------------------------------------------------- #
def _family_block(platform, media_key, spots):
    """逐區列（子公司樣式）：六區各 spots 檔。"""
    rows = [sm.CueRow(kind="main", region=r, spots=spots) for r in sc.REGIONS_ORDER]
    return sm.CueBlock(platform=platform, rows=rows)


def _carrefour_block(mag_spots, sup_spots):
    return sm.CueBlock(platform="家樂福", rows=[
        sm.CueRow(kind="main", region="量販", spots=mag_spots),
        sm.CueRow(kind="super", region="超市", spots=sup_spots),
    ])


@pytest.fixture
def ref_blocks():
    return [
        _family_block("全家廣播", "全家廣播", 140),
        _family_block("新鮮視", "新鮮視", 147),
        _carrefour_block(112, 192),
    ]


def test_reference_vector_exact(ref_blocks):
    """§1.5：三平台 + 合計逐一相等。"""
    r = sr.compute_reach(ref_blocks, sheetdata(), ndays=14)
    fam, tv, cf = r["by_platform"]

    assert fam["impressions"] == 629580
    assert rhu(fam["traffic"]) == 4372083
    assert tv["impressions"] == 492156
    assert rhu(tv["traffic"]) == 3417750
    assert cf["impressions"] == 860000
    assert rhu(cf["traffic"]) == 6480000
    assert cf["detail"]["量販"]["traffic"] == 5280000
    assert cf["detail"]["超市"]["traffic"] == 1200000

    assert r["total"]["impressions"] == 1981736
    assert rhu(r["total"]["traffic"]) == 14269833


def test_region_breakdown(ref_blocks):
    """逐地區明細：每區 曝光＝店數×檔次、Σ地區＝平台合計。"""
    r = sr.compute_reach(ref_blocks, sheetdata(), ndays=14)
    fam = r["by_platform"][0]
    regs = fam["regions"]
    assert len(regs) == 6                       # 六區
    # 北區 1673×140
    assert regs[0]["impressions"] == 1673 * 140
    # 各區曝光加總＝平台合計
    assert sum(x["impressions"] for x in regs) == fam["impressions"] == 629580
    assert sum(x["traffic"] for x in regs) == pytest.approx(fam["traffic"])
    # 家樂福：量販/超市兩列，曝光為固定常數 660000/200000
    cf = r["by_platform"][2]
    assert [x["impressions"] for x in cf["regions"]] == [660000, 200000]


def test_single_row_traffic_not_rounded():
    """§1.5：單列 Z 不取整。Z9 北區 1673×140、Z15 新鮮視北區 1209×147。"""
    d = sheetdata()
    north_fam = sm.CueBlock(platform="全家廣播",
                            rows=[sm.CueRow(kind="main", region="北區", spots=140)])
    north_tv = sm.CueBlock(platform="新鮮視",
                           rows=[sm.CueRow(kind="main", region="北區", spots=147)])
    z9 = sr.compute_reach([north_fam], d, 14)["by_platform"][0]["traffic"]
    z15 = sr.compute_reach([north_tv], d, 14)["by_platform"][0]["traffic"]
    assert z9 == pytest.approx(1626527.777, abs=1e-2)
    assert z15 == pytest.approx(1234187.5, abs=1e-6)


def test_client_lines_exact(ref_blocks):
    """§1.4/§6.3：客戶用三行文字逐字等於範本。"""
    r = sr.compute_reach(ref_blocks, sheetdata(), ndays=14)
    assert r["lines"] == [
        "【全家通路廣播總曝光次數 : 629580     曝光期間店舖總人流量 : 4372083 .】",
        "【TV總曝光次數 : 492156     曝光期間店舖總人流量 : 3417750 .】",
        "【萬家福．樂家康通路廣播總曝光次數 : 860000     曝光期間店舖總人流量 : 6480000 .】",
    ]


# --------------------------------------------------------------------------- #
# 萬家福除數開關（§6.1）
# --------------------------------------------------------------------------- #
def test_wjf_divisor_days_switch(monkeypatch):
    d = sheetdata()
    blk = _carrefour_block(112, 192)
    # days 模式、ndays=14：量販除數 14＝營業時數 14 → 與 hours 模式相同
    monkeypatch.setattr(sc, "REACH_WJF_DIVISOR", "days")
    r14 = sr.compute_reach([blk], d, ndays=14)["by_platform"][0]
    assert r14["detail"]["量販"]["traffic"] == 5280000
    # ndays=28：量販人流減半
    r28 = sr.compute_reach([blk], d, ndays=28)["by_platform"][0]
    assert r28["detail"]["量販"]["traffic"] == 2640000


def test_wjf_hours_from_daypart_equals_constant():
    """§6.1：由 DAYPART 解析的營業時數＝常數表。"""
    assert sr._wjf_hours("量販") == sc.REACH_WJF["量販"]["hours"] == 14
    assert sr._wjf_hours("超市") == sc.REACH_WJF["超市"]["hours"] == 24


# --------------------------------------------------------------------------- #
# 全省店數＝六區加總（§3.3 / §6.1）
# --------------------------------------------------------------------------- #
def test_national_stores_equals_region_sum():
    assert sum(sc.STORES_FALLBACK["全家廣播"].values()) == 4497
    assert sum(sc.STORES_FALLBACK["新鮮視"].values()) == 3348


# --------------------------------------------------------------------------- #
# 代理商 ④ 2008 全家 25萬 30秒（§3.3 / §6.1）：回饋開關
# --------------------------------------------------------------------------- #
def test_agency_family_reach_bonus_switch():
    d = sheetdata()
    m = sm.build_model("ag_2008_fam", 250000, date(2026, 6, 8), date(2026, 6, 28), data=d)
    s30 = next(s for s in m.sheets if s.seconds == 30)
    fam = s30.reach["by_platform"][0]
    assert fam["stores"] == 4497
    # 預設含回饋：4497 × (480 + 72 + 221)
    assert fam["impressions"] == 4497 * (480 + 72 + 221)
    # 關閉回饋：4497 × 480
    off = sr.compute_reach(s30.blocks, d, len(s30.days), include_bonus=False)
    assert off["by_platform"][0]["impressions"] == 4497 * 480


def test_agency_wjf_reach_constant_impressions():
    """代理商⑤萬家福：曝光固定 860000（與檔次天數無關）。"""
    d = sheetdata()
    m = sm.build_model("ag_2008_wjf", 250000, date(2026, 7, 17), date(2026, 7, 30), data=d)
    for s in m.sheets:
        assert s.reach["by_platform"][0]["impressions"] == 860000


# --------------------------------------------------------------------------- #
# 子公司整合（經 build_model 全流程，§6.2 數字一致性基礎）
# --------------------------------------------------------------------------- #
def _all_cell_values(xlsx_bytes):
    import io
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(xlsx_bytes))
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str):
                    yield c.coordinate, c.value


# --------------------------------------------------------------------------- #
# Excel 開關（§4.3 / §6.4）
# --------------------------------------------------------------------------- #
def test_excel_reach_switch_off_zero_leak(monkeypatch):
    """SHOW_REACH_IN_CUE=False 時：Excel 內完全不出現「曝光」二字（零洩漏）。"""
    import simple_excel as se
    monkeypatch.setattr(sc, "SHOW_REACH_IN_CUE", False)
    m = sm.build_model("sub_qp_fv", 200000, date(2026, 9, 21), date(2026, 10, 4),
                       data=sheetdata())
    xlsx = se.render(m, formulas=False)
    assert not any("曝光" in v for _, v in _all_cell_values(xlsx))


def test_excel_reach_default_on():
    """預設 SHOW_REACH_IN_CUE=True（2026-09-24 老闆令定：七組合皆印費用區左側）。"""
    assert sc.SHOW_REACH_IN_CUE is True


def test_excel_reach_switch_on_writes_lines():
    """SHOW_REACH_IN_CUE=True（預設）：子公司 Excel A 欄印出客戶各平台文字。"""
    import simple_excel as se
    m = sm.build_model("sub_qp_fv", 200000, date(2026, 9, 21), date(2026, 10, 4),
                       data=sheetdata())
    xlsx = se.render(m, formulas=False)
    a_vals = [v for coord, v in _all_cell_values(xlsx) if coord.startswith("A")]
    for line in m.sheets[0].reach["lines"]:
        assert line in a_vals


@pytest.mark.parametrize("key,extra", [
    ("ag_2008_fam", {"client": "統一企業", "product": "統一 "}),
    ("ag_2008_wjf", {"client": "統一企業", "product": "統一 "}),
    ("ag_carat_fam", {"client": "統一", "product": "麥香", "today": date(2026, 8, 31)}),
    ("ag_carat_wjf", {"client": "統一", "product": "純喫茶", "today": date(2026, 8, 11)}),
])
def test_agency_excel_reach_lines_on(key, extra):
    """SHOW_REACH_IN_CUE=True（預設）：四組代理商 Excel 也在 A 欄印出各平台結論。"""
    import simple_excel as se
    m = sm.build_model(key, 250000, date(2026, 9, 7), date(2026, 9, 20),
                       data=sheetdata(), **extra)
    xlsx = se.render(m, formulas=False)
    a_vals = [v for coord, v in _all_cell_values(xlsx) if coord.startswith("A")]
    lines = m.sheets[0].reach["lines"]
    assert lines, f"{key} 無 reach 行"
    for line in lines:
        assert line in a_vals, f"{key} A 欄缺行：{line}"


def test_subsidiary_reach_structure():
    m = sm.build_model("sub_qp_fv", 250000, date(2026, 9, 21), date(2026, 10, 4),
                       data=sheetdata())
    reach = m.sheets[0].reach
    assert {e["platform"] for e in reach["by_platform"]} == {"全家廣播", "新鮮視"}
    assert reach["total"]["impressions"] > 0
    assert len(reach["lines"]) == 2
