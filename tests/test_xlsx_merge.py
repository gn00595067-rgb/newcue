# -*- coding: utf-8 -*-
"""多個 CUE Excel 合成一個活頁簿（年約季約存檔附件）：內容、合併儲存格、欄寬、Logo 不變。"""
import io
from datetime import date

import pytest
from openpyxl import load_workbook

from agency_excel import generate_agency_excel
from annual_quarter_cue import build_wave_agency_model
from fixed_projects import FIXED_PROJECTS, build_project_rows, project_total_list
from excel_renderer import render_excel_into
from xlsx_merge import merge_xlsx

COMBOS = [{"media": "全家廣播", "region": "全省", "seconds": 15}, {"media": "家樂福", "region": "全省", "seconds": 20}]


def _agency(agency, start, end):
    m = build_wave_agency_model(agency, COMBOS, start, end, [800, 300], [100000, 90000], client_name="客")
    return generate_agency_excel(m, date(2026, 10, 13))


def _general(fmt):
    p = next(iter(FIXED_PROJECTS.values()))
    rows = build_project_rows(p)
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    render_excel_into(fmt, p["start"], p["end"], "客", "", "品", rows, ["1.x"], p["budget"], 0, "王小明",
                      project_total_list(rows), wb=wb, sheet_prefix="")
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def _sheets(b):
    return load_workbook(io.BytesIO(b)).worksheets


@pytest.mark.parametrize("parts", [
    [("波段1", "2008傳媒", date(2027, 1, 4), date(2027, 1, 31)), ("波段2", "2008傳媒", date(2027, 3, 1), date(2027, 3, 31))],
    [("波段1", "凱絡", date(2027, 1, 4), date(2027, 1, 31)), ("波段2", "凱絡", date(2027, 3, 1), date(2027, 3, 31))],
])
def test_merge_agency_keeps_content(parts):
    srcs = [(p, _agency(a, s, e)) for p, a, s, e in parts]
    merged = _sheets(merge_xlsx(srcs))
    flat = [(p, ws) for p, b in srcs for ws in _sheets(b)]
    assert len(merged) == len(flat)
    names = [ws.title for ws in merged]
    assert len(set(names)) == len(names) and all(n.startswith("波段") for n in names)
    for (p, src), dst in zip(flat, merged):
        assert [[c.value for c in r] for r in src.iter_rows()] == [[c.value for c in r] for r in dst.iter_rows()]
        assert {str(r) for r in src.merged_cells.ranges} == {str(r) for r in dst.merged_cells.ranges}
        assert len(src._images) == len(dst._images)
        for k, d in src.column_dimensions.items():
            assert dst.column_dimensions[k].width == d.width
        a1 = src.cell(5, 5)
        assert dst.cell(5, 5).border.left.style == a1.border.left.style


@pytest.mark.parametrize("fmt", ["東吳", "聲活", "鉑霖"])
def test_merge_general_formats(fmt):
    b = _general(fmt)
    merged = _sheets(merge_xlsx([("波段1", b), ("波段2", b)]))
    assert [ws.title for ws in merged][:1] == ["波段1"] and len(merged) == 2
    src = _sheets(b)[0]
    assert [[c.value for c in r] for r in src.iter_rows()] == [[c.value for c in r] for r in merged[1].iter_rows()]
    assert merged[1].page_setup.orientation == src.page_setup.orientation


def test_merge_skips_empty():
    assert [ws.title for ws in _sheets(merge_xlsx([("波段1", None)]))] == ["空"]
