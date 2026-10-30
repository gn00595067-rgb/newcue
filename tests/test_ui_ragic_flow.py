# -*- coding: utf-8 -*-
"""Ragic 存檔流程的畫面測試（AppTest；假 Ragic，不連線寫入）。
需要讀 Google 試算表設定檔（與正式 app 相同）；讀不到時略過。"""
import io
import os
from datetime import date

import pytest
from openpyxl import load_workbook

pytest.importorskip("streamlit.testing.v1")
from streamlit.testing.v1 import AppTest  # noqa: E402

HARNESS = os.path.join(os.path.dirname(__file__), "ui_harness_ragic.py")
USER = {"email": "a@x.com", "name": "王小明", "groups": []}
COMBOS = [{"media": "全家廣播", "region": "全省", "seconds": 15}, {"media": "家樂福", "region": "全省", "seconds": 20}]


def _app(mode, **state):
    at = AppTest.from_file(HARNESS, default_timeout=600)
    at.session_state["ragic_user"] = USER
    at.session_state["cue_mode"] = mode
    at.session_state["simple_payment"] = "115.12.31"   # 簡易模式存檔必填（佔位 XX 會擋存檔）
    for k, v in state.items():
        at.session_state[k] = v
    at.run()
    if any("設定檔載入失敗" in e.value for e in at.error):
        pytest.skip("讀不到 Google 試算表設定檔")
    return at


def _ok(at):
    assert not at.exception, at.exception[0].value
    assert not [e.value for e in at.error], [e.value for e in at.error]


def _rec(at):
    return at.session_state["cue_record"]


def test_simple_save_update_dirty_reopen():
    at = _app("簡易模式")
    at.text_input(key="simple_client").input("客戶A").run()
    at.button(key="simple_save_btn").click().run()
    _ok(at)
    assert _rec(at)["snap"]["version"] == 1
    assert not any("尚未存檔" in w.value for w in at.warning)
    at.session_state["simple_budget2"] = 300000
    at.run()
    assert any("尚未存檔" in w.value for w in at.warning)                  # 改了沒存要提醒
    at.button(key="simple_save_btn").click().run()
    _ok(at)
    assert _rec(at)["snap"]["version"] == 2 and list(at.session_state["_fake_db"]) == ["100"]   # 更新同一筆
    assert any("預算" in c for c in _rec(at)["snap"]["history"][-1]["changes"])
    # 換模式、弄亂欄位後從「我的 CUE」開啟
    at.session_state["cue_mode"] = "一般CUE"
    at.session_state["simple_client"] = "zzz"
    at.run()
    at.button(key="mycue_open").click().run()
    _ok(at)
    assert at.session_state["cue_mode"] == "簡易模式"
    assert at.text_input(key="simple_client").value == "客戶A" and at.session_state["simple_budget2"] == 300000
    assert not any("尚未存檔" in w.value for w in at.warning)              # 剛開啟＝與存檔一致


def test_fixed_project_save():
    at = _app("專案CUE")
    at.button(key="fp_save_btn").click().run()
    _ok(at)
    assert _rec(at)["snap"]["waves"][0]["rows"][0]["platform"] == "家樂福"


def test_annual_save_merges_wave_files_and_restores():
    waves = [{"start": date(2027, 1, 4), "end": date(2027, 1, 31), "combo_spots": [800, 300], "combo_net": [100000, 90000]},
             {"start": date(2027, 3, 1), "end": date(2027, 3, 31), "combo_spots": [800, 0], "combo_net": [100000, 0]}]
    at = _app("年約季約細CUE", aq_combos=COMBOS, aq_waves=waves)
    at.text_input(key="aq_client").input("年約客戶").run()
    at.button(key="aq_save_btn").click().run()
    _ok(at)
    snap = _rec(at)["snap"]
    assert len(snap["waves"]) == 2 and sum(r["net"] for w in snap["waves"] for r in w["rows"]) == 290000
    xl = at.session_state["_fake_last_files"]["1015320"][1]
    assert load_workbook(io.BytesIO(xl)).sheetnames == ["波段1", "波段2"]   # 各波合成一檔
    at.session_state["cue_mode"] = "簡易模式"
    at.session_state["aq_waves"] = []
    at.run()
    at.button(key="mycue_open").click().run()
    _ok(at)
    assert at.text_input(key="aq_client").value == "年約客戶"
    assert at.number_input(key="aq_tbl_net_0_1").value == 90000


def test_annual_agency_wave_settings_persist():
    waves = [{"start": date(2027, 1, 4), "end": date(2027, 1, 31), "combo_spots": [800, 300], "combo_net": [100000, 90000]}]
    at = _app("年約季約細CUE", aq_combos=COMBOS, aq_waves=waves, agency_unlocked=True)
    at.radio(key="aq_format").set_value("凱絡").run()
    rem = next(t for t in at.text_area if t.key.startswith("aq_ag_rem_0_"))
    rem.input("自訂備註一行").run()
    at.text_input(key="aq_ag_pay_0").input("自訂請款").run()
    at.button(key="aq_save_btn").click().run()
    _ok(at)
    assert _rec(at)["snap"]["header"]["format"] == "凱絡"
    at.session_state["cue_mode"] = "簡易模式"
    at.session_state["aq_waves"] = []
    at.run()
    at.button(key="mycue_open").click().run()
    _ok(at)
    rem = next(t for t in at.text_area if t.key.startswith("aq_ag_rem_0_"))
    assert rem.value == "自訂備註一行" and at.text_input(key="aq_ag_pay_0").value == "自訂請款"


def test_general_cue_save_and_reopen():
    at = _app("一般CUE")
    at.button(key="gen_save_btn").click().run()
    _ok(at)
    snap = _rec(at)["snap"]
    assert sum(r["net"] for r in snap["waves"][0]["rows"]) == int(snap["header"]["final"])
    at.button(key="mycue_open").click().run()
    _ok(at)
    assert at.session_state["cue_mode"] == "一般CUE"


@pytest.mark.parametrize("sub,btn,mode", [("自訂代理商CUE", "ag_save_btn", "代理商CUE"),
                                          ("一鍵CUE（模板）", "agsimple_save_btn", "代理商一鍵")])
def test_agency_save(sub, btn, mode):
    at = _app("代理商CUE", agency_unlocked=True, agency_sub_mode=sub)
    at.button(key=btn).click().run()
    _ok(at)
    assert _rec(at)["snap"]["mode"] == mode


def test_not_logged_in_no_save_button():
    at = AppTest.from_file(HARNESS, default_timeout=600)
    at.session_state["cue_mode"] = "簡易模式"
    at.run()
    assert not any(b.key == "simple_save_btn" for b in at.button)


# ---------------------------------------------------------------- 第 3 期：檔次調整（加贈）
def _enable_bonus(at, prefix, spots=3):
    at.run()   # AppTest：存檔觸發 st.rerun 後的下一次操作可能沒送進去，先補跑一次（瀏覽器操作無此問題）
    cb = next(c for c in at.checkbox if str(c.key).startswith(f"{prefix}cb_en_"))
    skey = cb.key[len(f"{prefix}cb_en_"):]
    cb.check()
    at.number_input(key=f"{prefix}cb_spots_{skey}").set_value(spots)
    at.run()
    return skey


def _xlsx_text(at):
    wb = load_workbook(io.BytesIO(at.session_state["_fake_last_files"]["1015320"][1]))
    return " ".join(str(c.value) for ws in wb.worksheets for r in ws.iter_rows() for c in r if c.value is not None)


def _totals(snap):
    rows = [r for w in snap["waves"] for r in w["rows"]]
    return sum(r["spots"] for r in rows), sum(r["net"] for r in rows), [r["kind"] for r in rows]


@pytest.mark.parametrize("mode,prefix,btn,state", [
    ("專案CUE", "fp_", "fp_save_btn", {}),
    ("簡易模式", "simple_", "simple_save_btn", {}),
    ("一般CUE", "", "gen_save_btn", {}),
    ("年約季約細CUE", "aq_", "aq_save_btn",
     {"aq_combos": COMBOS, "aq_waves": [{"start": date(2027, 1, 4), "end": date(2027, 1, 31),
                                         "combo_spots": [800, 300], "combo_net": [100000, 90000]}]}),
])
def test_bonus_adds_spots_keeps_net(mode, prefix, btn, state):
    at = _app(mode, **state)
    at.button(key=btn).click().run()
    _ok(at)
    s1, n1, _ = _totals(_rec(at)["snap"])
    _enable_bonus(at, prefix)
    assert any("尚未存檔" in w.value for w in at.warning)
    at.button(key=btn).click().run()
    _ok(at)
    snap = _rec(at)["snap"]
    s2, n2, kinds = _totals(snap)
    assert s2 > s1 and n2 == n1 and "bonus" in kinds                  # 加贈：檔次增加、實收不變
    assert snap["version"] == 2 and any("檔次" in c for c in snap["history"][-1]["changes"])
    assert "加贈" in _xlsx_text(at)                                    # Excel 重新產出含加贈列
    # 開啟還原後加贈設定仍在
    at.session_state["cue_mode"] = "專案CUE" if mode != "專案CUE" else "簡易模式"
    at.run()
    at.button(key="mycue_open").click().run()
    _ok(at)
    assert any(str(c.key).startswith(f"{prefix}cb_en_") and c.value for c in at.checkbox)
    assert not any("尚未存檔" in w.value for w in at.warning)


@pytest.mark.parametrize("fmt", ["2008傳媒", "凱絡"])
def test_bonus_agency_annual(fmt):
    waves = [{"start": date(2027, 1, 4), "end": date(2027, 1, 31), "combo_spots": [800, 300], "combo_net": [100000, 90000]}]
    at = _app("年約季約細CUE", aq_combos=COMBOS, aq_waves=waves, agency_unlocked=True)
    at.radio(key="aq_format").set_value(fmt).run()
    at.button(key="aq_save_btn").click().run()
    s1, n1, _ = _totals(_rec(at)["snap"])
    _enable_bonus(at, "aqag_", spots=4)
    at.button(key="aq_save_btn").click().run()
    _ok(at)
    s2, n2, _ = _totals(_rec(at)["snap"])
    assert s2 > s1 and n2 == n1


def test_bonus_agency_custom():
    at = _app("代理商CUE", agency_unlocked=True, agency_sub_mode="自訂代理商CUE")
    at.button(key="ag_save_btn").click().run()
    s1, n1, _ = _totals(_rec(at)["snap"])
    _enable_bonus(at, "ag_", spots=5)
    at.button(key="ag_save_btn").click().run()
    _ok(at)
    s2, n2, _ = _totals(_rec(at)["snap"])
    assert s2 == s1 + 5 * 28 and n2 == n1          # 預設走期 8/5～9/1 共 28 天，每日 +5


# ---------------------------------------------------------------- 第 5 期：編號、先存再下載、送出即成交、營運加贈、管理、統計
import re  # noqa: E402

CUSTOMER = {"rid": "9", "編號": "C01399", "名稱": "客戶甲股份有限公司", "統編": "12345678", "產品": "產品乙",
            "類別": "直客", "業務": "王小明", "主管": "李主管", "客服": ["陳客服"], "付款條件": "月結30天",
            "現折": "5", "退佣": "0"}
BOSS = {"email": "boss@x.com", "name": "營運主管 Boss", "groups": []}


def _after_rerun(at, *keys, **values):
    """AppTest：按鈕觸發 st.rerun 後，舊畫面元件可能已不存在，先補值再重跑（瀏覽器無此問題）。"""
    for k in keys:
        if k not in at.session_state:
            at.session_state[k] = ""
    for k, v in values.items():
        if k not in at.session_state:
            at.session_state[k] = v
    at.run()


def _has(at, key):
    return any(b.key == key for b in at.button)


def _apply_cue(at, prefix):
    """兩段式：先存檔（才有模擬編號）→ 🎫 申請 CUE號（已有 CUE號 時跳過）。"""
    at.run()
    if not _has(at, f"{prefix}_cue_apply") and not _has(at, f"{prefix}_cue_drop") and _has(at, f"{prefix}_btn"):
        at.button(key=f"{prefix}_btn").click().run()
        at.run()
    if _has(at, f"{prefix}_cue_apply"):
        at.button(key=f"{prefix}_cue_apply").click().run()
        at.run()


def _book(at, prefix):
    _apply_cue(at, prefix)
    at.date_input(key=f"{prefix}_od_pay").set_value(date(2027, 2, 28)).run()
    btn = at.button(key=f"{prefix}_submit")
    assert not btn.disabled, btn.help
    btn.click().run()
    _after_rerun(at, f"{prefix}_au_reason")
    _ok(at)


def test_save_number_and_download_after_save():
    at = _app("簡易模式", simple_client="客戶甲")
    assert not at.get("download_button")                            # 未存檔不能下載
    at.button(key="simple_save_btn").click().run()
    at.run()
    _ok(at)
    snap = _rec(at)["snap"]
    assert re.fullmatch(r"A\d{4}-001", snap["save_no"])            # 王小明無英文名 → 信箱 a@x.com 的 A
    assert snap["owner"]["email"] == "a@x.com"
    assert at.get("download_button")                                # 存檔後可下載（AppTest 讀不到下載鈕 key，只看有無）
    at.session_state["simple_budget2"] = 300000
    at.run()
    assert not at.get("download_button")                            # 改了沒存又不能下載
    at.button(key="simple_save_btn").click().run()
    at.run()
    assert _rec(at)["snap"]["save_no"] == snap["save_no"]           # 改版本沿用編號
    at.radio(key="simple_save_how").set_value(at.radio(key="simple_save_how").options[1]).run()
    at.button(key="simple_save_btn").click().run()
    at.run()
    assert _rec(at)["snap"]["save_no"].endswith("-002")             # 另存新的一筆給新號


def test_book_is_deal_then_adjust_updates_booking():
    at = _app("簡易模式", ragic_customer=CUSTOMER, simple_client="客戶甲", _fake_sign=True)
    _book(at, "simple_save")
    snap = _rec(at)["snap"]
    assert snap["status"] == "已訂檔" and snap["booking"]["cue_id"] == "1160101"
    assert at.session_state["_fake_booking_calls"][-1][0] == "submit"
    assert at.session_state["_fake_booking_serial"] == "S" + _rec(at)["rid"]   # 新存檔先存拿模擬編號再申請 CUE號
    assert len(at.session_state["_fake_db"]) == 1                              # 先存＋送出後再存＝同一筆
    assert len(at.session_state["_fake_booking_calls"][-1][2]) == 2          # 訂檔CUE表＋客戶回簽檔
    assert at.session_state["_fake_files"][_rec(at)["rid"]]["1015884"][0] == "回簽.pdf"
    _enable_bonus(at, "simple_")
    at.button(key="simple_save_btn").click().run()
    at.run()
    assert _rec(at)["snap"]["status"] == "已訂檔－有異動未送審"
    at.session_state["_fake_sign"] = False                                   # 沿用上次回簽檔
    _book(at, "simple_save")
    assert at.session_state["_fake_booking_calls"][-1] == ("update", "1160101", ["500"])
    assert _rec(at)["snap"]["status"] == "已訂檔"


def test_book_blocked_without_customer_signback_or_flag():
    at = _app("簡易模式", simple_client="客戶甲")
    btn = at.button(key="simple_save_submit")
    assert btn.disabled and "客戶回簽檔" in btn.help and "客戶編號" in btn.help
    at = _app("簡易模式", ragic_customer=CUSTOMER, simple_client="客戶甲", _fake_sign=True, _fake_booking_on=False)
    at.date_input(key="simple_save_od_pay").set_value(date(2027, 2, 28)).run()
    btn = at.button(key="simple_save_submit")
    assert btn.disabled and "RAGIC_BOOKING_ENABLED" in btn.help


def test_sales_cannot_apply_cue_or_book():
    """10-30：開放業務測試期間，申請 CUE號／送出訂檔限主管與開發者，業務只能存檔。"""
    at = _app("簡易模式", ragic_customer=CUSTOMER, simple_client="客戶甲", _fake_sign=True, _fake_can_book=False)
    at.button(key="simple_save_btn").click().run()
    at.run()
    _ok(at)
    assert _rec(at)["rid"]                                                        # 存檔照常
    btn = at.button(key="simple_save_cue_apply")
    assert btn.disabled and "限主管／開發者操作" in btn.help
    btn = at.button(key="simple_save_submit")
    assert btn.disabled and "限主管／開發者操作" in btn.help
    assert "_fake_booking_calls" not in at.session_state                         # 沒有任何 CUE號／訂檔呼叫


def test_order_panel_inline_customer_and_fields():
    """宜歆 10/21：缺客戶編號等找不到去哪填 → 訂檔區直接搜尋帶入客戶、補客戶類別，播證名稱預設廣告主-產品。"""
    at = _app("簡易模式", _fake_customers=[dict(CUSTOMER, 類別="")], _fake_sign=True)
    btn = at.button(key="simple_save_submit")
    assert btn.disabled and "客戶編號（① 客戶：搜尋帶入）" in btn.help
    at.text_input(key="simple_save_od_cust_q").input("客戶甲").run()
    at.button(key="simple_save_od_cust_apply").click().run()
    at.run()
    _ok(at)
    assert at.session_state["ragic_customer"]["編號"] == "C01399"
    assert at.session_state["simple_client"] == "客戶甲股份有限公司"            # 上方空白的客戶名稱回填
    btn = at.button(key="simple_save_submit")
    assert "客戶類別（請到 Ragic 客戶主檔補" in btn.help and "客戶編號" not in btn.help   # 連動欄：去主檔補
    at.session_state["ragic_customer"] = dict(CUSTOMER, 類別="廣代")                   # 主檔補好後重新帶入
    at.run()
    assert "素材後補" in at.text_input(key="simple_save_od_need").value                # 沒放素材 → 預設寫素材後補
    btn = at.button(key="simple_save_submit")
    assert btn.disabled and "CUE號（在 ⓪ 先申請" in btn.help                     # 兩段式：先申請 CUE號
    _apply_cue(at, "simple_save")                                              # 存檔 → 申請（AppTest 重跑會清欄位，故先做）
    assert at.session_state["_fake_booking_calls"][0] == ("apply", "S" + _rec(at)["rid"])   # 帶存檔模擬編號
    at.text_area(key="simple_save_od_mat_names").input("B篇.30秒").run()             # 只有篇名、素材後補
    at.toggle(key="simple_save_od_preview").set_value(True).run()                  # 仿 Ragic 預覽
    _ok(at)
    at.text_input(key="simple_save_od_need").input("10/21 提供CUE+回簽").run()
    btn = at.button(key="simple_save_submit")
    assert not btn.disabled, btn.help                                          # 收款日期取付款兌現日 115.12.31
    btn.click().run()
    _after_rerun(at, "simple_save_au_reason")
    _ok(at)
    import cue_store as cs
    p = cs.from_jsonable(_rec(at)["snap"]["review"]["params"])
    assert p["客戶類別"] == "廣代" and p["客服"] == ["陳客服"]
    # 只有篇名的素材 Ragic 會丟（10-22 實測 852）→ 併進需求說明，不建素材列
    assert p["需求說明"] == "10/21 提供CUE+回簽（素材後補：B篇.30秒）" and p["素材"] == []
    assert p["現折"] == "0.05" and p["文案"] == ["無"]


def test_ops_bonus_mark_and_audit_reject():
    at = _app("專案CUE", ragic_customer=CUSTOMER, fp_client="客戶甲", _fake_sign=True)
    _book(at, "fp_save")
    no = _rec(at)["snap"]["save_no"]
    # 營運主管開啟業務的存檔 → 加贈 → 存檔：標記營運加贈，編號、存檔人不變
    at.session_state["ragic_user"] = BOSS
    at.session_state["is_supervisor"] = True
    at.session_state["cue_mode"] = "CUE 管理"
    at.run()
    at.button(key="adm_open").click().run()
    _after_rerun(at)
    _enable_bonus(at, "fp_")
    at.button(key="fp_save_btn").click().run()
    _after_rerun(at, "fp_save_au_reason")
    snap = _rec(at)["snap"]
    assert snap["ops"]["by"] == BOSS["name"] and snap["save_no"] == no and snap["owner"]["email"] == "a@x.com"
    assert snap["history"][-1]["changes"][0] == "【營運加贈】"
    # 每次送出都要主管審核：送出更新後回到待審，主管退回
    _book(at, "fp_save")
    assert _rec(at)["snap"]["audit"]["result"] == "待審"
    at.text_input(key="fp_save_au_reason").input("金額不對").run()
    at.button(key="fp_save_au_no").click().run()
    _after_rerun(at, "fp_save_au_reason")
    snap = _rec(at)["snap"]
    assert snap["status"] == "退回" and snap["audit"]["result"] == "退回" and snap["audit"]["reason"] == "金額不對"


def test_manage_and_stats_views():
    at = _app("簡易模式", simple_client="客戶甲", ragic_customer=CUSTOMER, _fake_sign=True)
    at.button(key="simple_save_btn").click().run()
    at.run()
    _book(at, "simple_save")
    at.session_state["cue_mode"] = "專案CUE"
    at.run()
    at.button(key="fp_save_btn").click().run()
    at.run()
    # 業務：CUE 管理只看自己的；送出訂檔可直接從清單送
    at.session_state["cue_mode"] = "CUE 管理"
    at.run()
    _ok(at)
    df = at.dataframe[0].value
    assert len(df) == 2 and any(str(x).startswith("★") for x in df["編號"])
    # 統計
    at.session_state["cue_mode"] = "統計"
    at.run()
    _ok(at)
    m = {x.label: x.value for x in at.metric}
    assert m["出 CUE 張數"] == "2" and m["成交張數"] == "1" and m["成交率"] == "50%"
    # 主管：CUE 管理預設看待審，選一份直接核准
    import cue_store as cs
    at.session_state["ragic_user"] = BOSS
    at.session_state["is_supervisor"] = True
    at.session_state["cue_mode"] = "CUE 管理"
    at.run()
    _ok(at)
    assert at.selectbox(key="adm_status").value == "待審" and len(at.dataframe[0].value) == 1
    at.button(key="adm_au_ok").click().run()
    _after_rerun(at, "adm_au_reason")
    _ok(at)
    rid = next(r for r, rec in at.session_state["_fake_db"].items() if "1160101" in rec["1015319"])
    assert cs.parse_snapshot(at.session_state["_fake_db"][rid]["1015319"])["audit"]["result"] == "核准"
    assert not at.dataframe                                              # 待審清單清空
    # 其他業務看不到
    at.session_state["is_supervisor"] = False
    at.session_state["ragic_user"] = {"email": "b@x.com", "name": "B 業務 Bob", "groups": []}
    at.session_state["cue_mode"] = "CUE 管理"
    at.run()
    assert not at.dataframe


def test_manage_submit_saved_record():
    at = _app("專案CUE", ragic_customer=CUSTOMER, fp_client="客戶甲")
    at.button(key="fp_save_btn").click().run()
    at.run()
    at.session_state["cue_mode"] = "CUE 管理"
    at.session_state["_fake_sign"] = True
    at.run()
    at.date_input(key="adm_od_pay").set_value(date(2027, 2, 28)).run()
    assert at.button(key="adm_submit").disabled                              # 尚未申請 CUE號
    at.button(key="adm_cue_apply").click().run()
    _after_rerun(at, adm_od_company="", adm_od_cat="商品", adm_od_pay=None, adm_od_inv=None,
                 adm_od_cs=[], adm_od_cert="", adm_od_need="", adm_od_disc="0", adm_od_rebate="0", adm_od_mat="",
                 adm_od_cw=["無"], adm_od_mat_names="", adm_od_preview=False, adm_cue_exist="")
    at.date_input(key="adm_od_pay").set_value(date(2027, 2, 28)).run()
    btn = at.button(key="adm_submit")
    assert not btn.disabled, btn.help
    btn.click().run()
    _after_rerun(at, adm_od_company="", adm_od_cat="商品", adm_od_pay=None, adm_od_inv=None, adm_od_ccat="",
                 adm_od_cs=[], adm_od_cert="", adm_od_need="", adm_od_disc="0", adm_od_rebate="0", adm_od_mat="",
                 adm_od_cw=["無"], adm_od_mat_names="", adm_od_preview=False)
    _ok(at)
    rid = next(iter(at.session_state["_fake_db"]))
    import cue_store as cs
    snap = cs.parse_snapshot(at.session_state["_fake_db"][rid]["1015319"])
    assert snap["status"] == "已訂檔" and snap["booking"]["cue_id"] == "1160101"
