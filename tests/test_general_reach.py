# -*- coding: utf-8 -*-
"""一般 CUE（東吳／聲活／鉑霖）預估曝光／人流：rows → 文字，並寫入 Excel 費用區 A 欄。"""
import io
import os
import sys
from datetime import date

import pytest
from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import simple_config as sc
from simple_reach import compute_reach_from_rows
from excel_renderer import generate_excel_from_scratch
from html_generator import generate_html_preview


def _row(media, region, stores, daily, days, **kw):
    r = {"media": media, "region": region, "program_num": stores, "daypart": "07-23",
         "seconds": 20, "spots": daily * days, "schedule": [daily] * days,
         "rate_display": 1000, "pkg_display": 1000, "is_pkg_member": False}
    r.update(kw)
    return r


def _rows(days=14):
    return [
        _row("全家廣播", "北區", 1673, 10, days),
        _row("全家廣播", "中區", 850, 10, days),
        # 加贈列：program_num 為文字，店數沿用同平台同區
        _row("全家廣播", "北區", "加贈檔次", 2, days, pkg_display="加贈", is_custom_bonus=True),
        _row("新鮮視", "北區", 1209, 15, days),
        _row("家樂福", "全省量販", 68, 4, days),
        _row("家樂福", "全省超市", 250, 2, days, rate_display="計量販", pkg_display="計量販"),
    ]


def test_family_tv_carrefour_numbers():
    rows = _rows(14)
    reach = compute_reach_from_rows(rows, 14)
    plat = {e["platform"]: e for e in reach["by_platform"]}
    assert list(plat) == ["全家廣播", "新鮮視", "家樂福"]
    # 全家：北區 1673×(140+28) + 中區 850×140
    assert plat["全家廣播"]["impressions"] == 1673 * 168 + 850 * 140
    assert plat["全家廣播"]["traffic"] == pytest.approx(plat["全家廣播"]["impressions"] * sc.TRAFFIC_FACTOR)
    assert plat["新鮮視"]["impressions"] == 1209 * 210
    # 萬家福．樂家康：曝光固定、人流依量販 56 檔／超市 28 檔
    assert plat["家樂福"]["impressions"] == 860000
    assert plat["家樂福"]["detail"]["量販"]["spots"] == 56
    assert plat["家樂福"]["detail"]["超市"]["spots"] == 28
    assert len(reach["lines"]) == 3
    assert reach["lines"][0].startswith("【全家通路廣播總曝光次數 : ")


def test_uses_schedule_not_full_period_spots():
    """分月切表時 schedule 已切片、spots 仍是全期 → 以 schedule 為準。"""
    r = _row("全家廣播", "北區", 100, 10, 30)
    r["schedule"] = r["schedule"][:10]
    assert compute_reach_from_rows([r], 10)["by_platform"][0]["impressions"] == 100 * 100


def test_skips_absent_platforms():
    reach = compute_reach_from_rows([_row("全家廣播", "北區", 100, 1, 7)], 7)
    assert [e["platform"] for e in reach["by_platform"]] == ["全家廣播"]
    assert compute_reach_from_rows([], 7)["lines"] == []


def _col_a_texts(xlsx):
    wb = load_workbook(io.BytesIO(xlsx))
    return [[c.value for c in ws["A"] if isinstance(c.value, str) and "總曝光次數" in c.value]
            for ws in wb.worksheets]


@pytest.mark.parametrize("fmt", ["東吳", "聲活", "鉑霖"])
def test_excel_writes_reach_lines(fmt):
    days = 14
    s, e = date(2026, 11, 2), date(2026, 11, 15)
    rows = _rows(days)
    xlsx = generate_excel_from_scratch(fmt, s, e, "客戶", "12345678", "產品", rows,
                                       ["1.測試"], 250000, 0, "承辦人", 300000)
    expect = compute_reach_from_rows(rows, days)["lines"]
    assert _col_a_texts(xlsx) == [expect]


@pytest.mark.parametrize("fmt", ["東吳", "聲活"])
def test_excel_multi_month_each_sheet_own_reach(fmt):
    s, e = date(2026, 10, 20), date(2026, 11, 30)   # 42 天跨月 → 分兩張表
    days = (e - s).days + 1
    rows = [_row("全家廣播", "北區", 100, 1, days)]
    xlsx = generate_excel_from_scratch(fmt, s, e, "客戶", "", "產品", rows,
                                       ["1.測試"], 100000, 0, "承辦人", 100000)
    per_sheet = _col_a_texts(xlsx)
    assert len(per_sheet) == 2
    assert "總曝光次數 : 1,200 " in per_sheet[0][0]    # 10/20–10/31：12 天
    assert "總曝光次數 : 3,000 " in per_sheet[1][0]    # 11 月：30 天


def test_html_preview_shows_reach():
    rows = _rows(14)
    html = generate_html_preview(rows, 14, date(2026, 11, 2), date(2026, 11, 15), "客戶", "", "產品",
                                 "東吳", ["1.測試"], 300000, 262500, 250000, 0)
    html = "".join(html) if isinstance(html, list) else html
    assert "全家通路廣播總曝光次數" in html and "萬家福．樂家康通路廣播總曝光次數" in html


@pytest.mark.parametrize("fmt", ["東吳", "聲活", "鉑霖"])
def test_custom_bonus_row_renders(fmt):
    """業務加贈列（program_num=「加贈檔次」）不可讓聲活／鉑霖當掉；原樣顯示、不計入總店數。"""
    rows = _rows(14)
    xlsx = generate_excel_from_scratch(fmt, date(2026, 11, 2), date(2026, 11, 15), "客戶", "", "產品", rows,
                                       ["1.測試"], 250000, 0, "承辦人", 300000)
    ws = load_workbook(io.BytesIO(xlsx)).worksheets[0]
    col_c = [c.value for c in ws["C"]]
    assert "加贈檔次" in col_c
    if fmt != "東吳":
        # 總店數＝數字列加總（1673+850+1209+68+250），不含加贈列
        assert 1673 + 850 + 1209 + 68 + 250 in col_c


@pytest.mark.parametrize("fmt", ["東吳", "聲活", "鉑霖"])
def test_carrefour_rows_split_station(fmt):
    """一般 CUE 的萬家福．樂家康逐列分：量販＝萬家福／量販店、超市＝樂家康／超市；其他平台照舊合併。"""
    rows = _rows(14)
    xlsx = generate_excel_from_scratch(fmt, date(2026, 11, 2), date(2026, 11, 15), "客戶", "", "產品", rows,
                                       ["1.測試"], 250000, 0, "承辦人", 300000)
    ws = load_workbook(io.BytesIO(xlsx)).worksheets[0]
    vals = {(c.row, c.column): c.value for row in ws.iter_rows() for c in row if c.value is not None}
    wjf = [r for (r, col), v in vals.items() if col == 1 and v == "萬家福"]
    lkk = [r for (r, col), v in vals.items() if col == 1 and v == "樂家康"]
    assert len(wjf) == 1 and len(lkk) == 1
    assert vals[(wjf[0], 2)] == "量販店" and vals[(lkk[0], 2)] == "超市"
    assert "萬家福．樂家康" not in [v for (r, col), v in vals.items() if col == 1]
    # 第一欄只有全家／新鮮視合併，萬家福、樂家康各自一格
    merged_a = [m for m in ws.merged_cells.ranges if m.min_col == 1 and m.max_col == 1 and m.min_row != m.max_row]
    assert all(not (m.min_row <= wjf[0] <= m.max_row) and not (m.min_row <= lkk[0] <= m.max_row) for m in merged_a)


def test_html_carrefour_split_station():
    html = generate_html_preview(_rows(14), 14, date(2026, 11, 2), date(2026, 11, 15), "客戶", "", "產品",
                                 "東吳", ["1.測試"], 300000, 262500, 250000, 0)
    html = "".join(html) if isinstance(html, list) else html
    assert "<td>萬家福</td><td>量販店</td>" in html and "<td>樂家康</td><td>超市</td>" in html
