# -*- coding: utf-8 -*-
"""一般 CUE 固定專案（萬家福/樂家康 116年度中元限定專案）＋ 一般 CUE 備註改版。"""
import io
import os
import sys
from datetime import date

import pytest
from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from fixed_projects import FIXED_PROJECTS, build_project_rows, project_days, project_total_list
from utils import get_remarks_text, remark_color, build_cue_filename
from excel_renderer import generate_excel_from_scratch
from html_generator import generate_html_preview

NAME = "萬家福/樂家康 116年度中元限定專案"
P = FIXED_PROJECTS[NAME]
STORES = {"家樂福_量販": 62, "家樂福_超市": 217}


def _rows():
    return build_project_rows(P, STORES)


def _rem():
    return get_remarks_text(date(2026, 11, 30), "2027年8月", date(2027, 9, 30), project_remark=P["remark"])


# ---------- 排程（對照參考檔「中元專案」工作表） ----------

def test_period_and_seconds_fixed():
    assert (P["start"], P["end"], P["seconds"], P["budget"]) == (date(2027, 7, 27), date(2027, 8, 16), 20, 300000)
    assert project_days(P) == 21


def test_rows_match_reference_sheet():
    rows = _rows()
    assert [r["region"] for r in rows] == ["全省量販", "全省超市", "全省量販", "全省超市", "全省超市"]
    assert [r["daypart"] for r in rows] == ["09-23 中元專區", "00-24 中元專區", "09-23", "00-24", "00-24"]
    assert [sum(r["schedule"]) for r in rows] == [168, 288, 294, 504, 63]
    assert all(len(r["schedule"]) == 21 and r["seconds"] == 20 for r in rows)
    # 超市專區最後 3 天（8/14–8/16）每日 12 檔
    assert rows[1]["schedule"][:18] == [14] * 18 and rows[1]["schedule"][18:] == [12] * 3
    assert [r["rate_display"] for r in rows] == [240000, "計量販", 210000, "計量販", 19250]
    assert project_total_list(rows) == 469250
    assert all(r["is_pkg_member"] and r["nat_pkg_display"] == 300000 for r in rows)
    assert [r["skip_store_total"] for r in rows] == [False, False, True, True, True]


def test_store_count_fallback_when_sheet_missing():
    rows = build_project_rows(P, {})
    assert [r["program_num"] for r in rows[:2]] == [62, 217]


# ---------- 備註 ----------

def test_general_remarks_1_to_6():
    rem = get_remarks_text(date(2026, 11, 30), "2月", date(2027, 3, 10))
    assert rem == [
        "1.請於 2026/11/30 (一) 中午12:00前 回簽及進單，方可順利上檔。",
        "2.通路店鋪數與開機率至少七成(以上)。每日因加盟數調整，或遇店舖年度季度改裝、設備維護升級及保修等狀況，會有一定幅度增減。",
        "3.託播方需於上檔前 5 個工作天，提供廣告帶(mp3)、影片/影像 1920x1080 (mp4)。",
        "4.結案資料以電子版提供為主，不另行提供紙本結案。雲端資料保留2個月（自提供日起算），敬請於期限內自行下載留存，逾期資料恕不保留。",
        "5.雙方同意費用請款月份 : 2月，如有修正必要，將另行E-Mail告知，並視為正式合約之一部分。",
        "6.付款兌現日期：116.03.10",
    ]


def test_project_remark_7_only_for_project():
    assert _rem()[-1] == "7.此為116年度中元專案，限定6席。"
    assert len(get_remarks_text(date(2026, 11, 30), "2月", date(2027, 3, 10))) == 6


def test_remarks_blank_dates():
    rem = get_remarks_text(None, "", None)
    assert "____/__/__ (__)" in rem[0] and rem[5].endswith("___.__.__")


def test_remark_color_by_content():
    rem = _rem()
    assert [remark_color(x) for x in rem] == ["FF0000", "000000", "FF0000", "000000", "000000", "000000", "0000FF"]
    assert remark_color(rem[5], blue_payment=True) == "0000FF"
    # 舊版手改備註：素材在第 4 點仍為紅字；舊第 2 點節目異動為黑字
    assert remark_color("4.託播方需於上檔前 5 個工作天，提供廣告帶(mp3)。") == "FF0000"
    assert remark_color("2.以上節目名稱如有異動，以上檔時節目名稱為主。") == "000000"


# ---------- 三種格式輸出 ----------

def _xlsx(fmt):
    return generate_excel_from_scratch(fmt, P["start"], P["end"], "客戶", "12345678", "產品", _rows(),
                                       _rem(), P["budget"], 0, "承辦人", 469250, medium_label=P["medium"])


def _cells(ws):
    return [c for row in ws.iter_rows() for c in row if c.value is not None]


@pytest.mark.parametrize("fmt", ["東吳", "聲活", "鉑霖"])
def test_excel_three_formats(fmt):
    wb = load_workbook(io.BytesIO(_xlsx(fmt)))
    assert len(wb.worksheets) == 1        # 21 天跨月不拆表
    ws = wb.worksheets[0]
    cells = _cells(ws)
    values = [c.value for c in cells]
    assert 300000 in values                # Package-cost 合併一格
    assert values.count("計量販") == 2
    assert 1317 in values                  # 總檔次 168+288+294+504+63
    assert "09-23 中元專區" in values and "00-24 中元專區" in values
    # 第 7 點只出現一次、藍字
    c7 = [c for c in cells if isinstance(c.value, str) and c.value.startswith("7.此為116年度中元專案")]
    assert len(c7) == 1 and c7[0].font.color.rgb.endswith("0000FF")
    c1 = next(c for c in cells if isinstance(c.value, str) and c.value.startswith("1.請於"))
    assert c1.font.color.rgb.endswith("FF0000")
    if fmt == "東吳":
        assert P["medium"] in values
    else:
        assert 62 + 217 in [c.value for c in ws["C"]]   # 總店數不重複加


def test_html_preview_project():
    html = generate_html_preview(_rows(), 21, P["start"], P["end"], "客戶", "", "20秒 產品", "東吳", _rem(),
                                 469250, 315000, 300000, 0, medium_label=P["medium"])
    html = "".join(html) if isinstance(html, list) else html
    assert P["medium"] in html and "$300,000" in html and "7.此為116年度中元專案" in html
    assert "rowspan='5'" in html


def test_project_filename():
    fn = build_cue_filename("萬國通路", _rows(), 300000, "宜", today=date(2026, 10, 8), seg_label=P["short"])
    assert fn == "1008 萬國通路 萬家福.樂家康(20秒) 30萬專案-中元限定-宜.xlsx"


def test_general_excel_without_medium_label_unchanged():
    """不傳 medium_label 時 Medium 仍依平台自動組（一般排程不受影響）。"""
    xlsx = generate_excel_from_scratch("東吳", P["start"], P["end"], "客戶", "", "產品", _rows(),
                                       ["1.測試"], 300000, 0, "承辦人", 469250)
    ws = load_workbook(io.BytesIO(xlsx)).worksheets[0]
    assert "萬家福．樂家康" in [c.value for c in _cells(ws)]


# ---------- 過年限定專案（參考檔「過年專案」工作表） ----------

NY = FIXED_PROJECTS["萬家福/樂家康 116年度過年限定專案"]


def test_new_year_rows_match_reference_sheet():
    from fixed_projects import project_seconds_text
    assert (NY["start"], NY["end"], NY["budget"]) == (date(2027, 1, 21), date(2027, 2, 10), 300000)
    assert project_days(NY) == 21 and project_seconds_text(NY) == "15、20秒"
    rows = build_project_rows(NY, STORES)
    assert [r["seconds"] for r in rows] == [15, 20, 20, 20, 20, 20]
    assert [r["region"] for r in rows] == ["全省超市", "全省量販", "全省超市", "全省量販", "全省超市", "全省超市"]
    assert [r["daypart"] for r in rows] == ["00-24 賀歲拜年", "09-23 年貨大街專區", "00-24 年貨大街專區",
                                            "09-23", "00-24", "00-24"]
    assert [sum(r["schedule"]) for r in rows] == [126, 168, 288, 294, 504, 63]
    # 賀歲拜年 15 秒：220000/720×126×0.85
    assert [r["rate_display"] for r in rows] == [32725, 240000, "計量販", 210000, "計量販", 19250]
    assert project_total_list(rows) == 501975
    assert [r["skip_store_total"] for r in rows] == [False, False, True, True, True, True]


@pytest.mark.parametrize("fmt", ["東吳", "聲活", "鉑霖"])
def test_new_year_excel_three_formats(fmt):
    rows = build_project_rows(NY, STORES)
    rem = get_remarks_text(date(2026, 11, 30), "2027年2月", date(2027, 3, 31), project_remark=NY["remark"])
    xlsx = generate_excel_from_scratch(fmt, NY["start"], NY["end"], "客戶", "", "產品", rows, rem,
                                       NY["budget"], 0, "承辦人", 501975, medium_label=NY["medium"])
    wb = load_workbook(io.BytesIO(xlsx))
    assert len(wb.worksheets) == 1
    values = [c.value for c in _cells(wb.worksheets[0])]
    assert 300000 in values and 501975 in values and 1443 in values   # 總檔 126+168+288+294+504+63
    assert "7.此為116年度過年專案，限定6席。" in values
    assert "00-24 賀歲拜年" in values
    if fmt == "東吳":
        assert NY["medium"] in values and "15秒、20秒 產品" in values
    else:
        assert 62 + 217 in [c.value for c in wb.worksheets[0]["C"]]


def test_new_year_filename():
    fn = build_cue_filename("萬國通路", build_project_rows(NY, STORES), 300000, "宜",
                            today=date(2026, 10, 8), seg_label=NY["short"])
    assert fn == "1008 萬國通路 萬家福.樂家康(15.20秒) 30萬專案-過年限定-宜.xlsx"


# ---------- 簡易模式備註與一般 CUE 同一份 ----------

def test_simple_mode_remarks_same_as_general():
    import simple_model as sm
    import simple_config as sc
    combo = next(c for c in sc.COMBOS.values() if c.get("family", "subsidiary") == "subsidiary")
    out = sm._subsidiary_remarks(combo, date(2026, 11, 30), "2月", "116.03.10")
    assert [t for t, _ in out] == get_remarks_text(date(2026, 11, 30), "2月", date(2027, 3, 10))
    assert [red for _, red in out] == [True, False, True, False, False, False]


def test_payment_text_passthrough():
    assert get_remarks_text(None, "", "116.XX.XX")[5] == "6.付款兌現日期：116.XX.XX"
