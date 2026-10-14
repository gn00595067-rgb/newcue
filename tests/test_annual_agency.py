# -*- coding: utf-8 -*-
"""年約季約細CUE 套 2008／凱絡格式（docs/specs/格式選擇擴充.md 需求 2）。"""
import io
from datetime import date

import pytest
from openpyxl import load_workbook

import agency_cue as ac
from agency_excel import generate_agency_excel
from agency_ui import build_agency_html
from annual_quarter_cue import agency_combo_ok, build_wave_agency_model

COMBOS = [
    {"media": "全家廣播", "region": "全省", "seconds": 15},
    {"media": "全家廣播", "region": "全省", "seconds": 20},
    {"media": "家樂福", "region": "全省", "seconds": 20},
    {"media": "新鮮視", "region": "全省", "seconds": 15},     # 不支援 → 略過
    {"media": "全家廣播", "region": "北區", "seconds": 15},   # 分區 → 略過
]
S, E = date(2027, 1, 4), date(2027, 1, 31)


def test_combo_ok():
    assert [agency_combo_ok(c) for c in COMBOS] == [True, True, True, False, False]


@pytest.mark.parametrize("agency", ["2008傳媒", "凱絡"])
@pytest.mark.parametrize("comp", ac.COMP_OPTIONS)
def test_wave_model(agency, comp):
    spots = [1000, 600, 400, 500, 500]
    nets = [120000, 80000, 100000, 50000, 50000]
    m = build_wave_agency_model(agency, COMBOS, S, E, spots, nets, client_name="客", product_name="品",
                                comp_mode=comp, ac_pct=3.0)
    main = [s for s in m["sheets"] if s["budget"]]
    assert len(main) == 3                     # 新鮮視／分區不出
    assert m["budget_total"] == 300000
    for sh, sp, nt in zip(main, spots, nets):
        mrow = next(r for r in sh["rows"] if r["kind"] == ac.KIND_MAIN)
        if sh["platform"] == ac.PLATFORM_FAMILY:
            if agency == "2008傳媒":
                assert mrow["spots"] == sp and mrow["net_display"] == nt
                assert sh["fees"]["budget_net"] == nt
            else:
                assert mrow["net_display"] == nt
        else:
            assert mrow["net_display"] == nt
    wb = load_workbook(io.BytesIO(generate_agency_excel(m, date(2026, 10, 13))))
    assert len(wb.sheetnames) == len(m["sheets"])
    assert "客" in build_agency_html(m, date(2026, 10, 13))


def test_zero_spots_skipped():
    m = build_wave_agency_model("凱絡", COMBOS, S, E, [0, 0, 300, 0, 0], [0, 0, 90000, 0, 0])
    assert len(m["sheets"]) == 1 and m["sheets"][0]["platform"] == ac.PLATFORM_WJF


def _aq_rows():
    """全家全省（六區展開、共用 nat_pkg）＋萬家福（單列）＋新鮮視北區（單列）。"""
    rows = [{"combo_idx": 0, "media": "全家廣播", "region": r, "seconds": 15, "rate_display": 100000 + i,
             "pkg_display": 100000 + i, "is_pkg_member": True, "nat_pkg_display": 100000}
            for i, r in enumerate(["北區", "桃竹苗", "中區", "雲嘉南", "高屏", "東區"])]
    rows.append({"combo_idx": 1, "media": "家樂福", "region": "全省量販", "seconds": 20, "rate_display": 214285,
                 "pkg_display": 90000, "is_pkg_member": False})
    rows.append({"combo_idx": 2, "media": "新鮮視", "region": "北區", "seconds": 15, "rate_display": 50000,
                 "pkg_display": 10000, "is_pkg_member": False, "nat_pkg_display": 0})
    return rows


def test_wave_net_total_counts_national_once():
    from annual_quarter_cue import wave_net_total
    assert wave_net_total(_aq_rows()) == 200000       # 修正前會把六區定價加總（> 70 萬）


def test_wave_price_override_scales_combos():
    from annual_quarter_cue import apply_wave_price_override, wave_net_total
    rows = _aq_rows()
    assert apply_wave_price_override(rows, 150000) == 150000
    assert wave_net_total(rows) == 150000
    assert {r["nat_pkg_display"] for r in rows if r.get("is_pkg_member")} == {75000}


@pytest.mark.parametrize("agency", ["2008傳媒", "凱絡"])
def test_annual_default_no_extra_rows(agency):
    """宜歆 2026-10-14：年約季約不論內容都多一列 720 檔（自帶專案回饋）。年約檔次以業務填的為準：
    不另加凌晨補償、自帶回饋列，2008 萬家福也不加 10%。"""
    spots = [400, 0, 300, 0, 0]
    m = build_wave_agency_model(agency, COMBOS, S, E, spots, [100000, 0, 90000, 0, 0], ac_pct=3.0)
    for sh in m["sheets"]:
        kinds = [r["kind"] for r in sh["rows"]]
        assert ac.KIND_REBATE not in kinds and ac.KIND_SUPER_REBATE not in kinds and ac.KIND_COMP not in kinds
        mrow = next(r for r in sh["rows"] if r["kind"] == ac.KIND_MAIN)
        assert mrow["spots"] == (400 if sh["platform"] == ac.PLATFORM_FAMILY else 300)
