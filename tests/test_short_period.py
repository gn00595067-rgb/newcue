# -*- coding: utf-8 -*-
"""短天期／跨月拆表時某月只有少數幾天：三格式都不可當掉，且東吳乙方簽章區不可壓到甲方區。"""
import io
import os
import sys
from datetime import date, timedelta

import pytest
from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from excel_renderer import generate_excel_from_scratch
from utils import get_remarks_text


def _rows(days):
    return [{"media": "全家廣播", "region": "北區", "program_num": 1000, "daypart": "07-23", "seconds": 20,
             "spots": 10 * days, "schedule": [10] * days, "rate_display": 1000, "pkg_display": 1000,
             "is_pkg_member": False}]


def _gen(fmt, s, e):
    days = (e - s).days + 1
    return generate_excel_from_scratch(fmt, s, e, "客戶", "12345678", "產品", _rows(days),
                                       get_remarks_text(date(2026, 10, 10), "11月", date(2026, 12, 31)),
                                       500000, 0, "承辦人", 600000)


CASES = [
    (date(2026, 11, 1), date(2026, 11, 1)),     # 單月 1 天
    (date(2026, 11, 1), date(2026, 11, 4)),     # 單月 4 天
    (date(2026, 10, 31), date(2026, 12, 30)),   # 跨月拆表，10 月只有 1 天
    (date(2026, 10, 29), date(2026, 12, 2)),    # 頭尾月各 3／2 天
]


@pytest.mark.parametrize("fmt", ["東吳", "聲活", "鉑霖"])
@pytest.mark.parametrize("s,e", CASES)
def test_short_month_does_not_crash(fmt, s, e):
    wb = load_workbook(io.BytesIO(_gen(fmt, s, e)))
    assert wb.worksheets


@pytest.mark.parametrize("s,e", CASES)
def test_dongwu_party_b_block_right_of_party_a(s, e):
    for ws in load_workbook(io.BytesIO(_gen("東吳", s, e))).worksheets:
        cells = {c.value: c for row in ws.iter_rows() for c in row if isinstance(c.value, str)}
        a = next(c for v, c in cells.items() if v.startswith("甲    方："))
        b = next(c for v, c in cells.items() if v.startswith("乙    方："))
        assert a.row == b.row and b.column >= 8      # 甲方佔 A:G，乙方從 H 之後
