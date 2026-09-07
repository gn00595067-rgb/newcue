# -*- coding: utf-8 -*-
"""
簡易模式 smoke（§10.3）：七組合 × 多預算 × 多走期，全部能產出且結構正確。
"""
import io
from datetime import date, timedelta

import pytest
from openpyxl import load_workbook

import simple_model as sm
import simple_excel as se
import simple_config as sc
from fixtures_simple import sheetdata

BUDGETS = [100000, 125000, 250000, 333000, 1000000]
# (開始日, 天數)：含跨月 28 天（9/21 起）
PERIODS = [(date(2026, 9, 21), 7), (date(2026, 9, 21), 14),
           (date(2026, 9, 21), 28), (date(2026, 9, 21), 56)]
KEYS = sc.SUBSIDIARY_COMBOS + sc.AGENCY_COMBOS


@pytest.mark.parametrize("key", KEYS)
@pytest.mark.parametrize("budget", BUDGETS)
@pytest.mark.parametrize("start,days", PERIODS)
def test_smoke_all(key, budget, start, days):
    end = start + timedelta(days=days - 1)
    model = sm.build_model(key, budget, start, end, client="客戶", product="產品",
                           campaign="C", data=sheetdata())
    # 分頁數 = 秒數版本數
    assert len(model.sheets) == len(sc.COMBOS[key]["seconds"])

    for formulas in (True, False):
        b = se.render(model, formulas=formulas)
        wb = load_workbook(io.BytesIO(b))
        # 分頁名唯一、≤31 字
        names = wb.sheetnames
        assert len(names) == len(set(names)), f"{key} 分頁名重複"
        for n in names:
            assert len(n) <= 31, f"{key} 分頁名過長: {n}"
        assert len(wb.worksheets) == len(model.sheets)

    # 每張表：日期欄數 = 天數；每列 schedule 長度 = 天數
    for s in model.sheets:
        assert len(s.days) == days
        for blk in s.blocks:
            for r in blk.rows:
                if r.schedule is not None:
                    assert len(r.schedule) == days
                    assert sum(r.schedule) == r.spots     # 排檔加總 = 檔次
                    assert all(isinstance(x, int) for x in r.schedule)


@pytest.mark.parametrize("key", sc.SUBSIDIARY_COMBOS)
@pytest.mark.parametrize("budget", BUDGETS)
def test_smoke_subsidiary_budget_invariants(key, budget):
    model = sm.build_model(key, budget, date(2026, 9, 21), date(2026, 10, 4), data=sheetdata())
    for s in model.sheets:
        assert s.hidden_net_total <= budget + 1e-6
        assert s.hidden_net_total >= 0.99 * budget
        assert s.fees["grand"] == budget + s.fees["prod"] + s.fees["vat"]


def test_smoke_filenames():
    for key in KEYS:
        model = sm.build_model(key, 250000, date(2026, 9, 21), date(2026, 10, 4),
                               client="統一", sales="宜", data=sheetdata())
        assert model.filename.endswith(".xlsx")
        assert "萬專案" in model.filename
