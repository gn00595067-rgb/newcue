# -*- coding: utf-8 -*-
"""健康視CUE（全民健康視，診所螢幕）：單列、走期區間、無每日檔次（spec：docs/specs/健康視CUE.md）。
曝光次數＝384,000 × 預算 ÷ 該秒數專案價未稅；定價欄＝定價未稅 × 預算 ÷ 專案價未稅。
價格表讀 Google 設定表「HealthPricing」分頁，讀不到用下方內建值。"""
import io
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.styles import Alignment, Border, Font, Side

PLATFORM = "全民健康視"
BASE_IMPRESSIONS = 384000   # 每一個「專案價未稅」的曝光次數（各秒數相同）
VAT = 0.05
# 秒數 → (定價含稅, 300面直客報價未稅, 專案價含稅)；未稅＝含稅 ÷ 1.05（宜歆 10/29 #6）
DEFAULT_PRICING = {
    10: (262500, 200000, 175000),
    15: (393750, 300000, 262500),
    20: (525000, 400000, 315000),
    25: (656250, 500000, 420000),
    30: (787500, 600000, 472500),
}
FORMATS = ["東吳", "聲活", "鉑霖"]
BRANDS = {   # 抬頭與 FROM 公司（版面同參考檔，只換文字；鉑霖另加 Mobi 標誌）
    "東吳": ("東吳廣告-媒體計劃排程表 Media Schedule", "東吳廣告"),
    "聲活": ("聲活數位-媒體計劃排程表 Media Schedule", "聲活數位"),
    "鉑霖": ("鉑霖行動行銷-媒體計劃排程表 Mobi Media Schedule", "鉑霖行動行銷"),
}
DAYS_OPTIONS = ["週一至週五", "每日"]


def _rnd(x):
    return int(x + 0.5) if x >= 0 else -int(-x + 0.5)


def calc(seconds, budget, pricing=None):
    """回傳 {"impressions","list","net","vat","gross"}（皆整數）。"""
    list_tax, _direct, proj_tax = (pricing or DEFAULT_PRICING)[int(seconds)]
    proj_net = proj_tax / (1 + VAT)
    budget = float(budget or 0)
    net = _rnd(budget)
    vat = _rnd(budget * VAT)
    return {"impressions": _rnd(BASE_IMPRESSIONS * budget / proj_net),
            "list": _rnd(list_tax / (1 + VAT) * budget / proj_net),
            "net": net, "vat": vat, "gross": net + vat}


def roc_month(d):
    return f"{d.year - 1911}/{d.month}月"


def roc_date(d):
    return f"{d.year - 1911}/{d.month}/{d.day}"


def default_billing(end):
    """請款月份預設：走期結束的下一個月（參考檔 12/22 結束 → 116/1月）。"""
    nxt = date(end.year + (end.month == 12), end.month % 12 + 1, 1)
    return roc_month(nxt)


def default_payment(end):
    """付款兌現日預設：走期結束下個月的最後一天（同其他模式）。"""
    nxt = date(end.year + (end.month == 12), end.month % 12 + 1, 1)
    return date(nxt.year + (nxt.month == 12), nxt.month % 12 + 1, 1) - timedelta(days=1)


def remarks(sign_deadline, billing_month, payment_date):
    """健康視專用備註（不用 utils.get_remarks_text）。回傳 [(文字, 顏色)]，顏色 None＝黑。"""
    d = (f"{sign_deadline:%Y/%m/%d} ({'一二三四五六日'[sign_deadline.weekday()]})" if sign_deadline else "____/__/__ (__)")
    pay = roc_date(payment_date) if isinstance(payment_date, date) else (str(payment_date or "").strip() or "___/__/__")
    return [
        (f"1. 請於 {d} 下午1500前回簽。", "FF0000"),
        ("2. 因應場域屬性，診所開診與休診時段不一，故不保證每家診所每日與各時段曝光次數維持一致，曝光記錄以後台系統數據為最終依據。", None),
        ("3. 託播方需於上檔前 5 個工作天，提供影片/影像 1920x1080 (MP4)。", "FF0000"),
        (f"4. 雙方同意費用請款月份 : {billing_month} ，如有修正必要，將另行告知，並視為正式合約之一部分。", None),
        (f"5. 付款兌現日期：{pay}前兌現。", None),
        ("6. 若實際曝光次數未達購買總量，將延長執行期至達成保證曝光數為止。", "0000FF"),
    ]


def period_text(start, end, days_opt):
    return f"{start:%m/%d}-{end:%m/%d}\n執行{days_opt}"


def filename(client, seconds, budget, sales_nick="", ext="xlsx", today=None):
    today = today or date.today()
    parts = [today.strftime("%m%d"), client or "", f"{PLATFORM}({seconds}秒)", f"{(budget or 0) / 10000:g}萬"]
    from utils import safe_filename
    return safe_filename(" ".join(p for p in parts if p) + (f"-{sales_nick}" if sales_nick else "")) + f".{ext}"


# =============================================================================
# Excel（值版；版面照參考檔 1151005 全民健康視 1123-1222）
# =============================================================================
_F = "微軟正黑體"
_MED, _THIN = Side(style="medium"), Side(style="thin")
_C = Alignment(horizontal="center", vertical="center", wrap_text=True)
_L = Alignment(horizontal="left", vertical="center")
_R = Alignment(horizontal="right", vertical="center")
_MONEY = '"$"#,##0_);[Red]\\("$"#,##0\\)'
_NUM = "#,##0_);[Red](#,##0)"


def _font(size, bold=False, color=None):
    return Font(name=_F, size=size, bold=bold, color=color)


def build_excel(d, logo_bytes=None):
    """d：{"format","client","advertiser","tax_id","ad_name","sales","seconds","start","end","days","budget",
    "sign_deadline","billing_month","payment_date"}（pricing 可選）。回傳 xlsx bytes。"""
    r = calc(d["seconds"], d["budget"], d.get("pricing"))
    title, company = BRANDS.get(d.get("format"), BRANDS["鉑霖"])
    wb = Workbook()
    ws = wb.active
    ws.title = "全民健康視"
    for col, w in {"A": 1.8, "B": 32.8, "C": 24.9, "D": 21.2, "E": 27.9, "F": 48.8, "G": 47.3,
                   "H": 30.3, "I": 33.0, "J": 33.0}.items():
        ws.column_dimensions[col].width = w
    for row, h in {1: 54, 2: 60, 3: 41.5, 4: 41.5, 5: 52.45, 6: 41.5, 7: 71, 8: 71, 9: 123, 10: 87.75,
                   11: 87, 12: 87, 13: 87, 14: 40.5}.items():
        ws.row_dimensions[row].height = h

    def put(ref, val, font, align=None, fmt=None):
        c = ws[ref]
        c.value = val
        c.font = font
        if align:
            c.alignment = align
        if fmt:
            c.number_format = fmt
        return c

    put("B1", title, _font(36, True), _L)
    put("B2", "全民健康視專案", _font(30, True), _L)
    if logo_bytes:
        try:
            img = OpenpyxlImage(io.BytesIO(logo_bytes))
            scale = 150 / img.height
            img.height, img.width = 150, int(img.width * scale)
            img.anchor = "I1"
            ws.add_image(img)
        except Exception:  # noqa: BLE001
            pass
    put("B3", "TO：", _font(25, True), _L)
    put("C3", d.get("client", ""), _font(25, True), _L)
    put("B4", "FROM：", _font(25, True), _L)
    put("C4", f"{company} - {d.get('sales', '')}", _font(25, True), _L)
    put("B5", "客戶名稱：", _font(25, True), _L)
    put("C5", d.get("client", ""), _font(25, True), _L)
    put("B6", "廣告名稱：", _font(26, True), _L)
    put("C6", d.get("ad_name", ""), _font(26, True), _L)
    for row in (5, 6):
        for ci, col in enumerate("BCDEFGHIJ"):
            ws[f"{col}{row}"].border = Border(left=_MED if col == "B" else None, right=_MED if col == "J" else None,
                                              top=_MED if row == 5 else None, bottom=_MED)

    heads = ["頻道", "播出地區", "播出店數", "播出時間", "規格", "執行期間", "曝光次數", "定價", "專案價"]
    cols = "BCDEFGHIJ"
    for col, h in zip(cols, heads):
        ws.merge_cells(f"{col}7:{col}8")
        put(f"{col}7", h, _font(25, h == "專案價"), _C)
    vals = [f"{PLATFORM}\n全科別診所", "全台", "300+面", "07:00-22:00",
            f"{d['seconds']}秒\n影片/影像 1920x1080 (MP4)", period_text(d["start"], d["end"], d.get("days", DAYS_OPTIONS[0])),
            r["impressions"], r["list"], r["net"]]
    fmts = [None] * 6 + [_NUM, _MONEY, _MONEY]
    for col, v, f in zip(cols, vals, fmts):
        put(f"{col}9", v, _font(26, color="0000FF" if col == "J" else None), _C, f)
    put("F10", "Total", _font(25, True), _C)
    put("H10", r["impressions"], _font(26, True), _C, _NUM)
    put("I10", r["list"], _font(26, True), _C, _MONEY)
    put("J10", r["net"], _font(26, True, "0000FF"), _C, _MONEY)
    for row in (7, 8, 9, 10):
        for col in cols:
            ws[f"{col}{row}"].border = Border(left=_MED if col == "B" else _THIN, right=_MED if col == "J" else _THIN,
                                              top=_MED if row in (7, 9, 10) else None,
                                              bottom=_MED if row in (8, 10) else _THIN)

    put("B11", "廣  告  主：", _font(28, True), _L)
    put("C11", d.get("advertiser") or d.get("client", ""), _font(28), _L)
    put("B12", "統一編號：", _font(28, True), _L)
    put("C12", d.get("tax_id", ""), _font(28), _L)
    put("B13", "客戶簽章：", _font(28, True), _L)
    put("I11", "製作 ", _font(25, True), _R)
    put("J11", "客戶提供", _font(25, True), _C)
    put("I12", "5% VAT ", _font(25, True), _R)
    put("J12", r["vat"], _font(25, True), _C, _MONEY)
    put("I13", "專案含稅價 ", _font(26, True, "0000FF"), _C)
    put("J13", r["gross"], _font(30, True, "0000FF"), _C, _MONEY)
    for row in (11, 12, 13):
        for col in cols:
            ws[f"{col}{row}"].border = Border(left=_MED if col in ("B", "I") else None,
                                              right=_MED if col in ("H", "I", "J") else None,
                                              top=_THIN if col in ("I", "J") and row > 11 else None,
                                              bottom=_MED if row == 13 else (_THIN if col in ("I", "J") else None))

    put("B14", "Remarks：本排程表經雙方確認後視同合約之延伸，具同等法律約束力與效力", _font(26, True), _L)
    for i, (text, color) in enumerate(remarks(d.get("sign_deadline"), d.get("billing_month", ""), d.get("payment_date"))):
        ws.row_dimensions[15 + i].height = 34.5
        put(f"B{15 + i}", text, _font(24, True, color), _L)

    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = 9
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 1
    ws.print_area = "A1:J20"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# =============================================================================
# 畫面
# =============================================================================
def preview_html(d, pricing=None):
    """簡單預覽（數字與備註），完整版面看下載的 Excel／PDF。"""
    from utils import html_escape as e
    r = calc(d["seconds"], d["budget"], pricing)
    title, company = BRANDS.get(d.get("format"), BRANDS["鉑霖"])
    cells = [f"{PLATFORM}<br>全科別診所", "全台", "300+面", "07:00-22:00",
             f"{d['seconds']}秒<br>影片/影像 1920x1080 (MP4)", e(period_text(d["start"], d["end"], d["days"])).replace("\n", "<br>"),
             f"{r['impressions']:,}", f"${r['list']:,}", f"<b style='color:#00f'>${r['net']:,}</b>"]
    th = "".join(f"<th>{h}</th>" for h in ["頻道", "播出地區", "播出店數", "播出時間", "規格", "執行期間", "曝光次數", "定價", "專案價"])
    td = "".join(f"<td>{c}</td>" for c in cells)
    rem = "".join(f"<div style='color:#{c or '000'}'>{e(t)}</div>"
                  for t, c in remarks(d["sign_deadline"], d["billing_month"], d["payment_date"]))
    return f"""<div style="font-family:'Microsoft JhengHei',sans-serif;font-size:13px">
<div style="font-size:18px;font-weight:bold">{e(title)}</div><div style="font-weight:bold">全民健康視專案</div>
<div>TO：{e(d['client'])}　FROM：{e(company)} - {e(d['sales'])}　廣告名稱：{e(d.get('ad_name', ''))}</div>
<table border=1 cellspacing=0 cellpadding=4 style="border-collapse:collapse;text-align:center;margin:6px 0">
<tr>{th}</tr><tr>{td}</tr>
<tr><td colspan=5></td><td><b>Total</b></td><td><b>{r['impressions']:,}</b></td><td><b>${r['list']:,}</b></td><td><b style='color:#00f'>${r['net']:,}</b></td></tr>
<tr><td colspan=7 style="text-align:left">廣告主：{e(d.get('advertiser') or d['client'])}　統一編號：{e(d.get('tax_id', ''))}</td><td>製作</td><td>客戶提供</td></tr>
<tr><td colspan=7></td><td>5% VAT</td><td>${r['vat']:,}</td></tr>
<tr><td colspan=7></td><td style="color:#00f">專案含稅價</td><td style="color:#00f"><b>${r['gross']:,}</b></td></tr>
</table><div><b>Remarks：本排程表經雙方確認後視同合約之延伸，具同等法律約束力與效力</b></div>{rem}</div>"""


def render_health_cue(sales_map, pricing=None):
    import streamlit as st

    import cue_store as cs
    import ragic_ui
    from config import RAGIC_MAP
    from pdf_converter import get_cloud_logo_bytes, xlsx_bytes_to_pdf_bytes

    pricing = pricing or DEFAULT_PRICING
    ss = st.session_state
    ss.setdefault("hc_client", ss.get("temp_client_name", ""))
    ss.setdefault("hc_tax", ss.get("temp_tax_id", ""))
    fmt = st.radio("選擇格式", FORMATS, key="hc_format", horizontal=True)
    if not ragic_ui.customer_gate({"client": "hc_client", "tax": "hc_tax"}, key_prefix="cust"):
        return

    c1, c2, c3, c4 = st.columns(4)
    client = c1.text_input("客戶名稱", key="hc_client")
    tax = c1.text_input("統一編號", key="hc_tax")
    advertiser = c2.text_input("廣告主（空白＝同客戶名稱）", key="hc_advertiser")
    ad_name = c2.text_input("廣告名稱（選填）", key="hc_ad_name")
    secs = sorted(pricing)
    seconds = c3.selectbox("秒數", secs, index=secs.index(15) if 15 in secs else 0, key="hc_seconds",
                           format_func=lambda s: f"{s}秒")
    budget = c3.number_input("預算（未稅）", min_value=0, value=300000, step=10000, key="hc_budget")
    sales_options = list(sales_map.keys()) if sales_map else []
    if sales_options and "hc_sales" not in ss:
        ss["hc_sales"] = sales_options[ragic_ui.default_sales_index(sales_options) or 0]   # 預設登入者
    sales = c4.selectbox("業務名稱", sales_options, key="hc_sales") if sales_options else ""
    days = c4.radio("執行日", DAYS_OPTIONS, key="hc_days", horizontal=True)

    d1, d2 = st.columns(2)
    start = d1.date_input("開始日", date.today() + timedelta(days=14), key="hc_start")
    end = d2.date_input("結束日", start + timedelta(days=29), key="hc_end")
    if end < start:
        st.error("結束日不能早於開始日")
        return
    with st.expander("📝 備註欄位設定", expanded=True):
        r1, r2, r3 = st.columns(3)
        sign = r1.date_input("回簽期限", date.today() + timedelta(days=3), key="hc_sign_deadline")
        billing = r2.text_input("請款月份", default_billing(end), key="hc_billing_month")
        pay = r3.date_input("付款兌現日期（預設走期結束下月最後一天）", default_payment(end), key="hc_payment_date")
    if not budget:
        st.info("請輸入預算")
        return

    nick = sales_map.get(sales, sales) if sales else ""
    d = {"format": fmt, "client": client, "advertiser": advertiser, "tax_id": tax, "ad_name": ad_name,
         "sales": nick, "seconds": seconds, "start": start, "end": end, "days": days, "budget": budget,
         "sign_deadline": sign, "billing_month": billing, "payment_date": pay, "pricing": pricing}
    r = calc(seconds, budget, pricing)
    m1, m2, m3, m4 = st.columns(4)
    m1.metric("曝光次數", f"{r['impressions']:,}")
    m2.metric("定價", f"${r['list']:,}")
    m3.metric("專案價（未稅）", f"${r['net']:,}")
    m4.metric("專案含稅價", f"${r['gross']:,}")
    st.caption(f"曝光次數 ＝ 384,000 × 預算 ÷ {seconds}秒專案價未稅 {pricing[seconds][2] / 1.05:,.0f}")
    st.markdown("---")
    st.subheader("📥 檔案下載區")
    st.components.v1.html(preview_html(d, pricing), height=480, scrolling=True)

    def files():
        xl = build_excel(d, get_cloud_logo_bytes() if fmt == "鉑霖" else None)
        pdf, _, err = xlsx_bytes_to_pdf_bytes(xl)
        if not pdf:
            st.warning(f"PDF 生成失敗: {err}")
        return {"xlsx": (filename(client, seconds, budget, nick), xl),
                "pdf": (filename(client, seconds, budget, nick, ext="pdf"), pdf) if pdf else None}

    row = cs._row(PLATFORM, "全台", seconds, 0, None, net=r["net"], lst=r["list"])
    snap = cs.make_snapshot(
        "健康視CUE",
        {"client": client, "tax_id": tax, "product": ad_name, "format": fmt, "sales": nick, "start": start, "end": end,
         "budget": budget, "final": budget, "prod_cost": 0, "payment_date": pay, "impressions": r["impressions"],
         "advertiser": advertiser, "days": days},
        [{"label": "", "start": start, "end": end, "rows": [row]}],
        restore=ragic_ui.capture(["hc_format", "hc_client", "hc_tax", "hc_advertiser", "hc_ad_name", "hc_seconds",
                                  "hc_budget", "hc_sales", "hc_days", "hc_start", "hc_end", "hc_sign_deadline",
                                  "hc_billing_month", "hc_payment_date"]),
        customer=ss.get(ragic_ui.SS_CUSTOMER))
    ragic_ui.render_save_panel(
        snap, files=files, allow_style_excel=False, booking=False,
        extra_fields={RAGIC_MAP["date_sign"]: str(sign), RAGIC_MAP["bill_month"]: billing,
                      RAGIC_MAP["date_pay"]: str(pay)},
        key_prefix="hc_save")
