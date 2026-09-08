# -*- coding: utf-8 -*-
"""
純 Python xlsx → PDF 渲染器（openpyxl 讀取 + reportlab 向量繪製）。

用途：取代依賴伺服器端 LibreOffice 的轉檔，讓 Streamlit Cloud（無 apt/LibreOffice）
也能產出與真 Excel 高度擬真的 CUE PDF。輸入為「已排版完成、值版（無公式）」的 xlsx bytes。

支援：欄寬/列高幾何、合併格、框線（hair/thin/medium/thick/double）、實色填色、
字型（內嵌繁中 Noto Sans TC）、對齊/自動換行/縮小字型、佈景/索引色解析、
數字/日期格式（見 xlsx_numfmt）、A3/A4 橫直向、fit-to-page 縮放與垂直分頁、
print_area、列標題重複（print_title_rows）。
"""
import datetime
import io
import os
import re

import openpyxl
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.styles.colors import COLOR_INDEX

from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.lib.utils import ImageReader

_EMU_PT = 12700.0  # EMU per point

import xlsx_numfmt

_HERE = os.path.dirname(os.path.abspath(__file__))
_FONT_DIR = os.path.join(_HERE, "assets", "fonts")

FONT_REG = "CJK"
FONT_BOLD = "CJK-Bold"
_FONTS_READY = False


def _ensure_fonts():
    global _FONTS_READY
    if _FONTS_READY:
        return
    reg = os.path.join(_FONT_DIR, "NotoSansTC-Regular.ttf")
    bold = os.path.join(_FONT_DIR, "NotoSansTC-Bold.ttf")
    has_bold = os.path.exists(bold)
    pdfmetrics.registerFont(TTFont(FONT_REG, reg))
    pdfmetrics.registerFont(TTFont(FONT_BOLD, bold if has_bold else reg))
    # 防呆：Regular/Bold 的 PostScript face name 若相同，reportlab 會去重成同一字型、
    # 粗體會失效（曾因兩檔 PS name 都是 NotoSansTC-Thin 而整份 PDF 無粗體）。
    if has_bold:
        reg_face = pdfmetrics.getFont(FONT_REG).face.name
        bold_face = pdfmetrics.getFont(FONT_BOLD).face.name
        if reg_face == bold_face:
            raise RuntimeError(
                f"字型 PS name 重複({reg_face})，粗體會失效；"
                "請跑 tools/fix_font_names.py 修正 name table。")
    _FONTS_READY = True


# ---- 幾何換算 -----------------------------------------------------------
_MDW = 7.0            # Calibri 11 之數字寬（px），Excel 欄寬單位基準
_PX2PT = 72.0 / 96.0  # 像素 → 點


def _col_width_pt(width_chars):
    px = int(round(width_chars * _MDW)) + 5
    return px * _PX2PT


# ---- 顏色解析 -----------------------------------------------------------
# Excel 儲存格 theme 索引 → 佈景 clrScheme 次序（前兩對 dk/lt 互換）
_THEME_MAP = [1, 0, 3, 2, 4, 5, 6, 7, 8, 9, 10, 11]
_DEFAULT_SCHEME = ["000000", "FFFFFF", "44546A", "E7E6E6",
                   "4472C4", "ED7D31", "A5A5A5", "FFC000",
                   "5B9BD5", "70AD47", "0563C1", "954F72"]


def _parse_theme(wb):
    try:
        raw = wb.loaded_theme
        if not raw:
            return list(_DEFAULT_SCHEME)
        if isinstance(raw, bytes):
            raw = raw.decode("utf-8", "ignore")
        m = re.search(r"<a:clrScheme.*?</a:clrScheme>", raw, re.S)
        if not m:
            return list(_DEFAULT_SCHEME)
        names = ["dk1", "lt1", "dk2", "lt2", "accent1", "accent2",
                 "accent3", "accent4", "accent5", "accent6", "hlink", "folHlink"]
        scheme = []
        seg = m.group(0)
        for nm in names:
            mm = re.search(r"<a:%s>(.*?)</a:%s>" % (nm, nm), seg, re.S)
            hexv = "000000"
            if mm:
                inner = mm.group(1)
                c = re.search(r'srgbClr val="([0-9A-Fa-f]{6})"', inner)
                if c:
                    hexv = c.group(1).upper()
                else:
                    c = re.search(r'lastClr="([0-9A-Fa-f]{6})"', inner)
                    if c:
                        hexv = c.group(1).upper()
                    else:
                        c = re.search(r'sysClr val="(\w+)"', inner)
                        hexv = "FFFFFF" if (c and c.group(1) == "window") else "000000"
            scheme.append(hexv)
        while len(scheme) < 12:
            scheme.append("000000")
        return scheme
    except Exception:
        return list(_DEFAULT_SCHEME)


def _apply_tint(hexv, tint):
    if not tint:
        return hexv
    r = int(hexv[0:2], 16); g = int(hexv[2:4], 16); b = int(hexv[4:6], 16)
    if tint < 0:
        f = 1.0 + tint
        r, g, b = r * f, g * f, b * f
    else:
        r = r * (1 - tint) + 255 * tint
        g = g * (1 - tint) + 255 * tint
        b = b * (1 - tint) + 255 * tint
    return "%02X%02X%02X" % (int(round(r)), int(round(g)), int(round(b)))


def _resolve_color(color, scheme, default=None):
    """openpyxl Color → 'RRGGBB' 或 default。"""
    if color is None:
        return default
    try:
        t = color.type
    except Exception:
        return default
    if t == "rgb":
        rgb = color.rgb
        if not isinstance(rgb, str):
            return default
        if len(rgb) == 8:
            # ARGB；alpha 00 視為未設（openpyxl 預設值）→ 交給 default
            if rgb[:2] == "00" and rgb[2:] == "000000":
                return default
            return rgb[2:].upper()
        return rgb.upper()
    if t == "theme":
        try:
            idx = int(color.theme)
        except Exception:
            return default
        if 0 <= idx < len(_THEME_MAP):
            base = scheme[_THEME_MAP[idx]]
        else:
            base = "000000"
        return _apply_tint(base, getattr(color, "tint", 0.0) or 0.0)
    if t == "indexed":
        try:
            idx = int(color.indexed)
            hexv = COLOR_INDEX[idx]
            return hexv[2:].upper() if len(hexv) == 8 else hexv.upper()
        except Exception:
            return default
    return default


def _hex_to_rl(hexv):
    return (int(hexv[0:2], 16) / 255.0, int(hexv[2:4], 16) / 255.0, int(hexv[4:6], 16) / 255.0)


# ---- 框線寬度 -----------------------------------------------------------
_BORDER_PT = {
    # hair 調到 0.5（原 0.35 在 1.6× 預覽只有半像素、抗鋸齒後像虛線『連不起來』）
    "hair": 0.5, "thin": 0.75, "medium": 1.75, "thick": 2.75,
    "dotted": 0.5, "dashed": 0.6, "double": 0.6,
    "mediumDashed": 1.5, "mediumDashDot": 1.5, "dashDot": 0.75,
    "slantDashDot": 1.0, "mediumDashDotDot": 1.5, "dashDotDot": 0.75,
}


def _border_pt(style):
    return _BORDER_PT.get(style, 0.75)


# ---- 頁面尺寸 -----------------------------------------------------------
from reportlab.lib.pagesizes import A3, A4, A5, letter, landscape

_PAPER = {8: A3, 9: A4, 11: A5, 1: letter, None: A4}


def _page_size(ps):
    base = _PAPER.get(ps.paperSize, A4)
    if (ps.orientation or "portrait") == "landscape":
        return landscape(base)
    return base


# ---- 主要渲染 -----------------------------------------------------------
class _SheetLayout:
    def __init__(self, ws, scheme):
        self.ws = ws
        self.scheme = scheme
        # 決定要繪製的範圍
        self.min_col, self.min_row, self.max_col, self.max_row = self._bounds()
        self.col_w = self._col_widths()
        self.row_h = self._row_heights()
        self.merge_lookup = self._merges()

    def _bounds(self):
        ws = self.ws
        if ws.print_area:
            pa = ws.print_area
            if isinstance(pa, (list, tuple)):
                pa = pa[0]
            pa = str(pa).split(",")[0].replace("$", "")
            if "!" in pa:
                pa = pa.split("!")[-1]
            c1, r1, c2, r2 = range_boundaries(pa)
            return c1, r1, c2, r2
        dim = ws.calculate_dimension()
        c1, r1, c2, r2 = range_boundaries(dim.replace("$", ""))
        return c1 or 1, r1 or 1, c2 or 1, r2 or 1

    def _col_widths(self):
        ws = self.ws
        default = (ws.sheet_format.defaultColWidth or 8.43) if ws.sheet_format else 8.43
        w = {}
        for c in range(self.min_col, self.max_col + 1):
            letter = get_column_letter(c)
            dim = ws.column_dimensions.get(letter)
            if dim and dim.hidden:
                w[c] = 0.0
            elif dim and dim.width:
                w[c] = _col_width_pt(dim.width)
            else:
                w[c] = _col_width_pt(default)
        return w

    def _row_heights(self):
        ws = self.ws
        default = (ws.sheet_format.defaultRowHeight or 15.0) if ws.sheet_format else 15.0
        h = {}
        for r in range(self.min_row, self.max_row + 1):
            dim = ws.row_dimensions.get(r)
            if dim and dim.hidden:
                h[r] = 0.0
            elif dim and dim.height is not None:
                h[r] = float(dim.height)
            else:
                h[r] = float(default)
        return h

    def _merges(self):
        """回傳 dict：(row,col) -> (r1,c1,r2,c2)。"""
        look = {}
        for rng in self.ws.merged_cells.ranges:
            c1, r1, c2, r2 = rng.min_col, rng.min_row, rng.max_col, rng.max_row
            for r in range(r1, r2 + 1):
                for c in range(c1, c2 + 1):
                    look[(r, c)] = (r1, c1, r2, c2)
        return look

    def content_width(self):
        return sum(self.col_w[c] for c in range(self.min_col, self.max_col + 1))

    def x_of(self, col):
        return sum(self.col_w[c] for c in range(self.min_col, col))

    def total_height(self, r_from, r_to):
        return sum(self.row_h[r] for r in range(r_from, r_to + 1))


def _wrap_text(text, font, size, max_w):
    """回傳行陣列。處理明確換行 + 依寬度自動換行（中文逐字、英數依空白）。"""
    lines = []
    for para in str(text).split("\n"):
        if not para:
            lines.append("")
            continue
        if stringWidth(para, font, size) <= max_w or max_w <= 0:
            lines.append(para)
            continue
        cur = ""
        i = 0
        while i < len(para):
            ch = para[i]
            trial = cur + ch
            if stringWidth(trial, font, size) > max_w and cur:
                # 嘗試在最後空白斷行（英數）
                sp = cur.rfind(" ")
                if sp > 0 and ord(ch) < 128:
                    lines.append(cur[:sp])
                    cur = cur[sp + 1:] + ch
                else:
                    lines.append(cur)
                    cur = ch
            else:
                cur = trial
            i += 1
        if cur:
            lines.append(cur)
    return lines


def _draw_sheet_page(c, lay, scheme, rows, title_rows, scale, page_w, page_h,
                     m_left, m_top, printable_w, printable_h, hcenter, vcenter):
    """把指定列（title_rows + rows）畫成一頁。"""
    ws = lay.ws
    draw_rows = list(title_rows) + list(rows)

    # 該頁內容高度、寬度（未縮放）
    content_w = lay.content_width()
    content_h = sum(lay.row_h[r] for r in draw_rows)
    used_w = content_w * scale
    used_h = content_h * scale

    off_x = m_left + ((printable_w - used_w) / 2 if hcenter else 0)
    off_y_top = page_h - m_top - ((printable_h - used_h) / 2 if vcenter else 0)

    c.saveState()
    c.translate(off_x, off_y_top)
    c.scale(scale, scale)
    # 之後以「y 向下為負」的座標系繪製；(0,0)=內容左上

    # 建立列 → y_top（未縮放，向下累加）對應
    y_top = {}
    acc = 0.0
    for r in draw_rows:
        y_top[r] = acc
        acc += lay.row_h[r]

    def cell_rect(r1, c1, r2, c2):
        x = lay.x_of(c1)
        w = sum(lay.col_w[cc] for cc in range(c1, c2 + 1))
        yt = y_top[r1]
        h = sum(lay.row_h[rr] for rr in range(r1, r2 + 1) if rr in y_top)
        return x, yt, w, h

    # 1) 先畫填色（避免蓋住框線）
    seen_anchor = set()
    for r in draw_rows:
        for col in range(lay.min_col, lay.max_col + 1):
            cell = ws.cell(row=r, column=col)
            mg = lay.merge_lookup.get((r, col))
            if mg:
                if (mg[0], mg[1]) in seen_anchor or (r, col) != (mg[0], mg[1]):
                    if (r, col) != (mg[0], mg[1]):
                        continue
                seen_anchor.add((mg[0], mg[1]))
                anchor = ws.cell(row=mg[0], column=mg[1])
                # 合併範圍在本頁可見的部分
                r2 = min(mg[2], draw_rows[-1])
                x, yt, w, h = cell_rect(mg[0] if mg[0] in y_top else r, mg[1],
                                        r2, mg[3])
                fill = anchor.fill
            else:
                x, yt, w, h = cell_rect(r, col, r, col)
                fill = cell.fill
            if fill is not None and fill.patternType == "solid":
                hexv = _resolve_color(fill.fgColor, scheme, None)
                if hexv and hexv != "FFFFFF":
                    c.setFillColorRGB(*_hex_to_rl(hexv))
                    c.rect(x, -(yt + h), w, h, stroke=0, fill=1)

    # 2) 文字
    for r in draw_rows:
        for col in range(lay.min_col, lay.max_col + 1):
            mg = lay.merge_lookup.get((r, col))
            if mg and (r, col) != (mg[0], mg[1]):
                continue
            cell = ws.cell(row=r, column=col)
            if cell.value is None or cell.value == "":
                continue
            if mg:
                r2 = min(mg[2], draw_rows[-1])
                x, yt, w, h = cell_rect(mg[0], mg[1], r2, mg[3])
            else:
                x, yt, w, h = cell_rect(r, col, r, col)
            _draw_text(c, cell, scheme, x, yt, w, h)

    # 3) 框線（畫在最上層；避開合併內部線）
    for r in draw_rows:
        for col in range(lay.min_col, lay.max_col + 1):
            cell = ws.cell(row=r, column=col)
            b = cell.border
            if not b:
                continue
            mg = lay.merge_lookup.get((r, col))
            x, yt, w, h = cell_rect(r, col, r, col)
            top = -yt
            bot = -(yt + h)
            left = x
            right = x + w
            edges = {"left": True, "right": True, "top": True, "bottom": True}
            if mg:
                edges["left"] = (col == mg[1])
                edges["right"] = (col == mg[3])
                edges["top"] = (r == mg[0])
                edges["bottom"] = (r == mg[2])
            if edges["top"]:
                _draw_border(c, b.top, scheme, left, top, right, top, horizontal=True)
            if edges["bottom"]:
                _draw_border(c, b.bottom, scheme, left, bot, right, bot, horizontal=True)
            if edges["left"]:
                _draw_border(c, b.left, scheme, left, top, left, bot, horizontal=False)
            if edges["right"]:
                _draw_border(c, b.right, scheme, right, top, right, bot, horizontal=False)

    # 4) 圖片（logo）
    _draw_images(c, lay, ws, y_top, draw_rows)

    c.restoreState()


def _draw_images(c, lay, ws, y_top, draw_rows):
    for img in getattr(ws, "_images", []):
        try:
            a = img.anchor
            frm = a._from
            col0 = frm.col          # 0-indexed
            row1 = frm.row + 1      # → 1-indexed
            if row1 not in y_top:
                continue            # 不在本頁
            x = lay.x_of(col0 + 1) + (frm.colOff or 0) / _EMU_PT
            yt = y_top[row1] + (frm.rowOff or 0) / _EMU_PT
            ext = getattr(a, "ext", None)
            to = getattr(a, "to", None)
            if ext is not None:
                w = ext.width / _EMU_PT
                h = ext.height / _EMU_PT
            elif to is not None:
                x2 = lay.x_of(to.col + 1) + (to.colOff or 0) / _EMU_PT
                y2 = (y_top.get(to.row + 1, yt) + (to.rowOff or 0) / _EMU_PT)
                w = max(x2 - x, 1)
                h = max(y2 - yt, 1)
            else:
                w = (img.width or 100) * _PX2PT
                h = (img.height or 100) * _PX2PT
            data = img._data() if hasattr(img, "_data") else img.ref
            reader = ImageReader(io.BytesIO(data) if isinstance(data, (bytes, bytearray)) else data)
            c.drawImage(reader, x, -(yt + h), w, h, mask="auto", preserveAspectRatio=False)
        except Exception:
            continue


def _draw_border(c, side, scheme, x1, y1, x2, y2, horizontal):
    if side is None or not side.style:
        return
    style = side.style
    hexv = _resolve_color(side.color, scheme, "000000") or "000000"
    c.setStrokeColorRGB(*_hex_to_rl(hexv))
    lw = _border_pt(style)
    c.setLineWidth(lw)
    dash = None
    if "Dash" in style or style in ("dashed", "dotted"):
        dash = [2, 1.5] if "dot" not in style.lower() else [1, 1]
    if dash:
        c.setDash(dash)
    if style == "double":
        gap = 1.2
        if horizontal:
            c.line(x1, y1 + gap / 2, x2, y2 + gap / 2)
            c.line(x1, y1 - gap / 2, x2, y2 - gap / 2)
        else:
            c.line(x1 + gap / 2, y1, x2 + gap / 2, y2)
            c.line(x1 - gap / 2, y1, x2 - gap / 2, y2)
    else:
        c.line(x1, y1, x2, y2)
    if dash:
        c.setDash([])


_PAD = 2.0  # 儲存格內距（點，未縮放）


def _draw_text(c, cell, scheme, x, yt, w, h):
    fmt = cell.number_format
    text, is_red = xlsx_numfmt.format_value(cell.value, fmt)
    if text is None or text == "":
        return
    font_obj = cell.font
    bold = bool(font_obj and font_obj.bold)
    size = float(font_obj.size) if (font_obj and font_obj.size) else 11.0
    fontname = FONT_BOLD if bold else FONT_REG

    # 顏色：紅字格式優先，其次字型色，預設黑
    if is_red:
        color = "FF0000"
    else:
        color = _resolve_color(font_obj.color if font_obj else None, scheme, "000000") or "000000"
    c.setFillColorRGB(*_hex_to_rl(color))

    al = cell.alignment
    halign = al.horizontal if al else None
    valign = al.vertical if al else None
    wrap = bool(al and al.wrap_text)
    shrink = bool(al and al.shrink_to_fit)

    avail_w = w - 2 * _PAD
    # general 對齊：數字靠右、其他靠左
    if halign in (None, "general"):
        halign = "right" if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool) else "left"

    # 會計格式 fill 分隔：符號靠左、數字靠右（單列置中）
    if "\x00" in text:
        left, right = text.split("\x00", 1)
        left = left.strip(); right = right.strip()
        c.setFont(fontname, size)
        by = -(yt + (h + size * 0.72) / 2)
        if left:
            c.drawString(x + _PAD, by, left)
        if right:
            rw = stringWidth(right, fontname, size)
            c.drawString(x + w - _PAD - rw, by, right)
        return

    if shrink and not wrap:
        while size > 4 and stringWidth(text, fontname, size) > avail_w:
            size -= 0.5

    if wrap or "\n" in str(text):
        lines = _wrap_text(text, fontname, size, avail_w)
    else:
        lines = [text]

    line_h = size * 1.18
    block_h = line_h * len(lines)
    # 垂直位置（y 向下為負）：cell 上緣 = -yt，下緣 = -(yt+h)
    if valign == "top":
        start_baseline = -(yt + _PAD + size)
    elif valign == "bottom":
        start_baseline = -(yt + h - _PAD - (block_h - size))
    else:  # center / None
        start_baseline = -(yt + (h - block_h) / 2 + size)

    for i, line in enumerate(lines):
        by = start_baseline - i * line_h
        lw = stringWidth(line, fontname, size)
        if halign == "center":
            lx = x + (w - lw) / 2
        elif halign == "right":
            lx = x + w - _PAD - lw
        else:
            lx = x + _PAD
        c.setFont(fontname, size)
        c.drawString(lx, by, line)


def _fit_scale(ps, lay, printable_w, printable_h):
    """依 fit-to-page 設定計算縮放比例與是否允許垂直分頁。"""
    pr = lay.ws.sheet_properties.pageSetUpPr
    fit = bool(pr is not None and pr.fitToPage)
    content_w = lay.content_width()
    if not fit:
        s = (ps.scale or 100) / 100.0
        return s, True
    fw = ps.fitToWidth
    fh = ps.fitToHeight
    # openpyxl 預設 fitToWidth/Height 皆為 1
    fw = 1 if fw is None else fw
    fh = 1 if fh is None else fh
    scales = []
    if fw and fw >= 1:
        scales.append(printable_w / content_w / fw)
    total_h = lay.total_height(lay.min_row, lay.max_row)
    if fh and fh >= 1:
        scales.append(printable_h / total_h / fh)
    if not scales:
        s = printable_w / content_w
    else:
        s = min(scales)
    s = min(s, 1.0)
    allow_paginate = not (fh and fh >= 1)   # fitH=0 → 允許垂直多頁
    return s, allow_paginate


def _paginate_rows(lay, rows, title_rows, scale, printable_h):
    """把資料列切成多頁（每頁高度 ≤ 可列印高度/scale）。"""
    cap = printable_h / scale
    title_h = sum(lay.row_h[r] for r in title_rows)
    pages = []
    cur = []
    cur_h = title_h
    for r in rows:
        rh = lay.row_h[r]
        if cur and cur_h + rh > cap:
            pages.append(cur)
            cur = [r]
            cur_h = title_h + rh
        else:
            cur.append(r)
            cur_h += rh
    if cur:
        pages.append(cur)
    return pages or [[]]


def _parse_hf(raw, sheet_name, ctx=None):
    """解析 Excel 頁首/頁尾字串（&-碼）→ (clean_text, size, bold)。
    ctx: {'file','date','time'} 供 &F/&D/&T 填值（缺則留空）。"""
    ctx = ctx or {}
    if not raw:
        return "", None, False
    out = []
    size = None
    bold = False
    i = 0
    n = len(raw)
    while i < n:
        ch = raw[i]
        if ch != "&":
            out.append(ch); i += 1; continue
        i += 1
        if i >= n:
            break
        nx = raw[i]
        if nx == "&":
            out.append("&"); i += 1; continue
        if nx == '"':                       # &"font,style"
            j = raw.find('"', i + 1)
            spec = raw[i + 1:j if j != -1 else n]
            if "bold" in spec.lower():
                bold = True
            i = (j + 1) if j != -1 else n; continue
        if nx.isdigit():                    # &<size>
            j = i
            while j < n and raw[j].isdigit():
                j += 1
            try:
                size = float(raw[i:j])
            except ValueError:
                pass
            i = j; continue
        if nx in "Bb":
            bold = True; i += 1; continue
        if nx == "K":                       # &Krrggbb 顏色（略）
            i += 1
            if raw[i:i + 6].strip():
                i += 6
            continue
        if nx in "IUESXYiuesxy":            # 斜體/底線/刪除線等：略
            i += 1; continue
        if nx in "Aa":
            out.append(sheet_name or ""); i += 1; continue
        if nx in "Ff":                      # &F 檔名
            out.append(ctx.get("file", "")); i += 1; continue
        if nx in "Dd":                      # &D 日期
            out.append(ctx.get("date", "")); i += 1; continue
        if nx in "Tt":                      # &T 時間
            out.append(ctx.get("time", "")); i += 1; continue
        if nx in "ZzPpNnGg":                # 路徑/頁碼/總頁/圖：留空
            i += 1; continue
        i += 1
    return "".join(out), size, bold


def _draw_hf_section(c, text, size, bold, x, y, align, avail_w):
    if not text:
        return
    font = FONT_BOLD if bold else FONT_REG
    sz = size or 10.0
    # 過寬則縮小以免溢出
    while sz > 5 and stringWidth(text, font, sz) > avail_w:
        sz -= 0.5
    c.setFont(font, sz)
    c.setFillColorRGB(0, 0, 0)
    w = stringWidth(text, font, sz)
    if align == "center":
        c.drawString(x - w / 2, y, text)
    elif align == "right":
        c.drawString(x - w, y, text)
    else:
        c.drawString(x, y, text)


_HF_BAND = 16.0   # 頁首/頁尾單行概估高度(pt，含上下間距)；供內容區讓位用


def _hf_visible(hf, sheet_name, ctx=None):
    """該頁首/頁尾任一段解析後是否有實際可見文字。"""
    for part in ("left", "center", "right"):
        text, _sz, _bold = _parse_hf(getattr(hf, part).text, sheet_name, ctx)
        if text.strip():
            return True
    return False


def _draw_header_footer(c, ws, page_w, page_h, m_left, m_right, pm, sheet_name, ctx=None):
    hmar = (pm.header or 0.3) * 72
    fmar = (pm.footer or 0.3) * 72
    printable_w = page_w - m_left - m_right
    cx = m_left + printable_w / 2
    rx = page_w - m_right
    for hf, is_header in ((ws.oddHeader, True), (ws.oddFooter, False)):
        y = (page_h - hmar) if is_header else fmar
        for part, align, x in (("left", "left", m_left),
                               ("center", "center", cx),
                               ("right", "right", rx)):
            raw = getattr(hf, part).text
            text, size, bold = _parse_hf(raw, sheet_name, ctx)
            _draw_hf_section(c, text, size, bold, x, y, align, printable_w)


def _title_rows(ws, min_row, max_row):
    tr = ws.print_title_rows
    if not tr:
        return []
    m = re.findall(r"\$?(\d+)\$?:\$?(\d+)", str(tr))
    rows = []
    for a, b in m:
        rows.extend(range(int(a), int(b) + 1))
    return [r for r in rows if min_row <= r <= max_row]


def render_xlsx_to_pdf(xlsx_bytes, filename="", made_date=None):
    """xlsx bytes → PDF bytes（純 Python，不需 LibreOffice）。
    filename：頁尾 &F 顯示的檔名（自動去 .xlsx）；made_date：頁尾 &D 日期(datetime/date)。"""
    _ensure_fonts()
    hf_ctx = {
        "file": (filename or "").rsplit(".xlsx", 1)[0],
        "date": made_date.strftime("%Y/%m/%d") if made_date else "",
        "time": made_date.strftime("%H:%M") if isinstance(made_date, datetime.datetime) else "",
    }
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=True)
    scheme = _parse_theme(wb)

    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    first = True

    sheets = [ws for ws in wb.worksheets if ws.sheet_state == "visible"] or wb.worksheets
    for ws in sheets:
        lay = _SheetLayout(ws, scheme)
        ps = ws.page_setup
        page_w, page_h = _page_size(ps)
        pm = ws.page_margins
        m_left = (pm.left or 0.2) * 72
        m_right = (pm.right or 0.2) * 72
        m_top = (pm.top or 0.2) * 72
        m_bottom = (pm.bottom or 0.2) * 72
        # 頁首/頁尾若有「實際會顯示」的文字，內容區需讓出該帶，否則會壓到內容
        # （常見於 bottom 邊界極小、footer 邊界較大時：頁尾落在內容區內重疊，如簽名列）。
        hmar = (pm.header or 0.3) * 72
        fmar = (pm.footer or 0.3) * 72
        if _hf_visible(ws.oddHeader, ws.title, hf_ctx):
            m_top = max(m_top, hmar + _HF_BAND)
        if _hf_visible(ws.oddFooter, ws.title, hf_ctx):
            m_bottom = max(m_bottom, fmar + _HF_BAND)
        printable_w = page_w - m_left - m_right
        printable_h = page_h - m_top - m_bottom

        scale, allow_paginate = _fit_scale(ps, lay, printable_w, printable_h)

        title_rows = _title_rows(ws, lay.min_row, lay.max_row)
        all_rows = list(range(lay.min_row, lay.max_row + 1))

        # 所有列依自然順序分頁；標題列（print_title_rows）在第一頁自然出現，
        # 第二頁起才於頁首重複。
        if allow_paginate:
            row_pages = _paginate_rows(lay, all_rows, [], scale, printable_h)
        else:
            row_pages = [all_rows]

        hcenter = bool(ws.print_options and ws.print_options.horizontalCentered)
        vcenter = bool(ws.print_options and ws.print_options.verticalCentered)

        for idx, pg in enumerate(row_pages):
            if not first:
                c.showPage()
            first = False
            c.setPageSize((page_w, page_h))
            repeat_titles = title_rows if (idx > 0 and title_rows and
                                           title_rows[-1] < pg[0]) else []
            _draw_sheet_page(c, lay, scheme, pg, repeat_titles, scale, page_w, page_h,
                             m_left, m_top, printable_w, printable_h, hcenter, vcenter)
            _draw_header_footer(c, ws, page_w, page_h, m_left, m_right, pm, ws.title, hf_ctx)

    c.showPage()
    c.save()
    return buf.getvalue()
