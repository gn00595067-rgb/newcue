# -*- coding: utf-8 -*-
"""
樣式母版產生器（§2.2）—— 從 cueexample 代理商範本第 1 分頁抽出「只有樣式、無值」的母版。

母版存到 assets/style_masters/*.xlsx，commit 進 repo；渲染時由 simple_style_master
把樣式整份轉印到新分頁再填值，代理商版面即與範本 100% 相同。

做法：載入範本 → 只留第 1 分頁 → 清掉所有儲存格值（樣式不動）→ 移除圖片 →
      移除 LAST+2 之後的欄寬設定（範本殘留幾百個無意義欄寬）→ 存檔。
母版內不得殘留任何值（測試 test_simple_smoke 會驗）。
"""
import os
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, column_index_from_string

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CUE = os.path.join(HERE, "cueexample")
OUT = os.path.join(HERE, "assets", "style_masters")

# key: (範本檔名, 分頁index, 最後日期欄LAST)
MASTERS = {
    "ag_2008_fam":  ("0902 2008-統一系列 全家單組 (10.15.20.30秒) 25萬專案.xlsx", 0, 29),
    "ag_2008_wjf":  ("0902 2008-統一系列 萬家福&樂家康單組  (10.15.20秒) 25萬專案.xlsx", 0, 38),
    "ag_carat_fam": ("0902 凱絡-統一系列 全家單組 (10.15.20.30秒) 25萬專案.xlsx", 0, 32),
    "ag_carat_wjf": ("0902 凱絡-統一系列 萬家福&樂家康單組  (10.15.20秒) 25萬專案.xlsx", 0, 21),
}


def build_one(key, fname, idx, last):
    wb = load_workbook(os.path.join(CUE, fname))
    keep = wb.worksheets[idx]
    for ws in list(wb.worksheets):
        if ws is not keep:
            wb.remove(ws)
    ws = keep
    # 清值（樣式保留）
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None:
                c.value = None
    # 移除圖片
    ws._images = []
    # 移除 LAST+2 之後的欄寬（範本殘留噪音）
    for letter in list(ws.column_dimensions.keys()):
        if column_index_from_string(letter) > last + 2:
            del ws.column_dimensions[letter]
    os.makedirs(OUT, exist_ok=True)
    path = os.path.join(OUT, f"{key}.xlsx")
    wb.save(path)
    # 驗證無殘值
    chk = load_workbook(path)
    vals = [c.value for r in chk.worksheets[0].iter_rows() for c in r if c.value is not None]
    print(f"{key}: saved {path}  merged={len(chk.worksheets[0].merged_cells.ranges)}  殘值={len(vals)}")
    assert not vals, f"{key} 母版仍有殘值：{vals[:5]}"


def main():
    for key, (fname, idx, last) in MASTERS.items():
        build_one(key, fname, idx, last)
    print("OK all masters built.")


if __name__ == "__main__":
    import sys
    sys.stdout.reconfigure(encoding="utf-8")
    main()
