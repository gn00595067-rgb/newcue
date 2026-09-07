# -*- coding: utf-8 -*-
"""
範本傾印工具 (Template Dumper) —— 簡易模式 v2 開發輔助（§2）

輸入 xlsx 路徑與分頁索引，印出每個非空儲存格的
  值 / 公式 / 字型大小 / 粗體 / 字色 / 底色 / 數字格式 / 對齊 / wrap / 四邊框線，
以及 merged ranges、欄寬、列高、列印設定（方向/紙張/fitToPage/邊距/列印範圍/標題列）。

用途：開發每張表之前先 dump 對應範本，渲染完再 dump 自己的輸出，逐格比對。

用法：
  python tools/dump_template.py "cueexample/xxx.xlsx"           # 列出分頁清單
  python tools/dump_template.py "cueexample/xxx.xlsx" 0         # dump 第 0 個分頁
  python tools/dump_template.py "cueexample/xxx.xlsx" 0 --max-row 40
"""
import argparse
import sys
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def _color(c):
    """把 openpyxl 顏色物件壓成短字串（rgb 或 theme）。"""
    if c is None:
        return "-"
    rgb = getattr(c, "rgb", None)
    if isinstance(rgb, str) and rgb not in ("00000000",):
        return rgb
    theme = getattr(c, "theme", None)
    if theme is not None:
        tint = getattr(c, "tint", 0) or 0
        return f"theme{theme}/{tint:.2f}"
    return "-"


def _fill(cell):
    f = cell.fill
    if f is None or f.patternType != "solid":
        return "-"
    return _color(f.fgColor)


def _border(cell):
    b = cell.border
    def s(side):
        return side.style if (side and side.style) else "."
    return f"T{s(b.top)} B{s(b.bottom)} L{s(b.left)} R{s(b.right)}"


def _align(cell):
    a = cell.alignment
    h = a.horizontal or "-"
    v = a.vertical or "-"
    w = "W" if a.wrap_text else "."
    return f"{h}/{v}/{w}"


def dump_sheet(ws, max_row=None, max_col=None):
    print(f"=== SHEET: {ws.title!r} ===")
    # 列印設定
    ps = ws.page_setup
    pm = ws.page_margins
    print("--- PRINT ---")
    print(f"orientation={ps.orientation} paperSize={ps.paperSize} "
          f"fitToPage={ws.sheet_properties.pageSetUpPr.fitToPage if ws.sheet_properties.pageSetUpPr else None} "
          f"fitW={ps.fitToWidth} fitH={ps.fitToHeight} scale={ps.scale}")
    print(f"print_area={ws.print_area} title_rows={ws.print_title_rows} title_cols={ws.print_title_cols}")
    print(f"margins L={pm.left} R={pm.right} T={pm.top} B={pm.bottom}")
    print(f"centerH={ws.print_options.horizontalCentered} centerV={ws.print_options.verticalCentered} "
          f"gridlines={ws.sheet_view.showGridLines} zoom={ws.sheet_view.zoomScale}")
    hf = ws.HeaderFooter
    try:
        print(f"footer L={hf.oddFooter.left.text!r} C={hf.oddFooter.center.text!r} R={hf.oddFooter.right.text!r}")
    except Exception:
        pass

    # 合併
    print("--- MERGED ---")
    print(sorted(str(r) for r in ws.merged_cells.ranges))

    # 欄寬
    print("--- COL WIDTH ---")
    widths = []
    for letter, dim in sorted(ws.column_dimensions.items()):
        if dim.width:
            widths.append(f"{letter}={dim.width:.1f}{'(H)' if dim.hidden else ''}")
    print(" ".join(widths))

    # 列高
    print("--- ROW HEIGHT ---")
    heights = []
    for idx, dim in sorted(ws.row_dimensions.items()):
        if dim.height:
            heights.append(f"{idx}={dim.height:.1f}")
    print(" ".join(heights))

    # 儲存格
    print("--- CELLS (row,col letter | value | formula-flag | sz bold color fill fmt align border) ---")
    mr = max_row or ws.max_row
    mc = max_col or ws.max_column
    for r in range(1, mr + 1):
        for c in range(1, mc + 1):
            cell = ws.cell(row=r, column=c)
            v = cell.value
            if v is None or v == "":
                continue
            fml = "F" if (isinstance(v, str) and v.startswith("=")) else "."
            fnt = cell.font
            sz = fnt.size
            bold = "B" if fnt.bold else "."
            col = _color(fnt.color)
            loc = f"{get_column_letter(c)}{r}"
            vs = repr(v)
            if len(vs) > 60:
                vs = vs[:57] + "...'"
            print(f"{loc:<5} {fml} {vs:<62} | sz{sz} {bold} fc={col} bg={_fill(cell)} "
                  f"nf={cell.number_format!r} al={_align(cell)} bd=[{_border(cell)}]")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("path")
    ap.add_argument("sheet", nargs="?", default=None, help="分頁索引 (0-based)；省略則列出分頁")
    ap.add_argument("--max-row", type=int, default=None)
    ap.add_argument("--max-col", type=int, default=None)
    args = ap.parse_args()

    wb = load_workbook(args.path, data_only=False)
    if args.sheet is None:
        print(f"FILE: {args.path}")
        for i, name in enumerate(wb.sheetnames):
            ws = wb[name]
            print(f"  [{i}] {name!r}  dims={ws.dimensions} maxRow={ws.max_row} maxCol={ws.max_column}")
        return
    idx = int(args.sheet)
    ws = wb.worksheets[idx]
    dump_sheet(ws, args.max_row, args.max_col)


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
