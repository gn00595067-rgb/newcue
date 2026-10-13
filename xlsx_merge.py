# -*- coding: utf-8 -*-
"""
把多個 Excel（bytes）合成一個活頁簿：逐張工作表複製值、樣式、合併儲存格、欄寬列高、列印設定與圖片（Logo）。
用途：年約季約存檔到 Ragic 時，各波 CUE 合成一個附件（每波一張，工作表名加「波段N」前綴）。
openpyxl 沒有跨活頁簿複製工作表，樣式物件需逐一 copy。
"""
import io
from copy import copy

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XLImage

_BAD = set(':\\/?*[]')


def _title(prefix, name, multi, used):
    base = f"{prefix}-{name}" if multi else prefix
    base = "".join(ch for ch in base if ch not in _BAD)[:31] or "Sheet"
    t, n = base, 2
    while t in used:
        suffix = f"({n})"
        t = base[:31 - len(suffix)] + suffix
        n += 1
    used.add(t)
    return t


def _copy_sheet(src, dst):
    # 先合併（之後對 MergedCell 設樣式才不會被 merge_cells 重設框線）
    for rng in src.merged_cells.ranges:
        dst.merge_cells(str(rng))
    for row in src.iter_rows():
        for c in row:
            d = dst.cell(row=c.row, column=c.column)
            if c.__class__.__name__ != "MergedCell":
                d.value = c.value
            if c.has_style:
                d.font = copy(c.font)
                d.border = copy(c.border)
                d.fill = copy(c.fill)
                d.number_format = c.number_format
                d.protection = copy(c.protection)
                d.alignment = copy(c.alignment)
    for k, dim in src.column_dimensions.items():
        dd = dst.column_dimensions[k]
        dd.width, dd.hidden = dim.width, dim.hidden
        if dim.min and dim.max:
            dd.min, dd.max = dim.min, dim.max
    for k, dim in src.row_dimensions.items():
        dst.row_dimensions[k].height = dim.height
        dst.row_dimensions[k].hidden = dim.hidden
    dst.sheet_format.defaultRowHeight = src.sheet_format.defaultRowHeight
    dst.sheet_view.showGridLines = src.sheet_view.showGridLines
    dst.sheet_view.zoomScale = src.sheet_view.zoomScale
    ps, dps = src.page_setup, dst.page_setup
    dps.orientation, dps.paperSize = ps.orientation, ps.paperSize
    dps.fitToWidth, dps.fitToHeight, dps.scale = ps.fitToWidth, ps.fitToHeight, ps.scale
    if src.sheet_properties.pageSetUpPr is not None:
        dst.sheet_properties.pageSetUpPr = copy(src.sheet_properties.pageSetUpPr)
    dst.print_options = copy(src.print_options)
    dst.page_margins = copy(src.page_margins)
    if src.print_area:
        dst.print_area = src.print_area
    if src.print_title_rows:
        dst.print_title_rows = src.print_title_rows
    dst.oddFooter.left.text = src.oddFooter.left.text
    dst.oddFooter.right.text = src.oddFooter.right.text
    for img in getattr(src, "_images", []):
        new = XLImage(io.BytesIO(img._data()))
        new.anchor = img.anchor
        new.width, new.height = img.width, img.height
        dst.add_image(new)


def merge_xlsx(parts):
    """parts：[(前綴, xlsx_bytes), ...] → 合併後 bytes。來源只有一張工作表時名稱＝前綴，多張則「前綴-原名」。"""
    out = Workbook()
    out.remove(out.active)
    used = set()
    for prefix, data in parts:
        if not data:
            continue
        wb = load_workbook(io.BytesIO(data))
        multi = len(wb.worksheets) > 1
        for ws in wb.worksheets:
            _copy_sheet(ws, out.create_sheet(_title(prefix, ws.title, multi, used)))
    if not out.worksheets:
        out.create_sheet("空")
    bio = io.BytesIO()
    out.save(bio)
    return bio.getvalue()
