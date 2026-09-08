# -*- coding: utf-8 -*-
"""
簡易模式 v2 代理商 Excel（版型 B 2008 / C 凱絡）—— 樣式母版版（§2）

版面（欄寬/列高/合併/框線/字級/底色/列印）由 simple_style_master 從 0902 範本轉印，
本檔只負責「填值」與「動態覆寫」（六日底色、月份、日期、星期、回饋文字、Logo）。
"""
import os

from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def _set_side(ws, r, c, **sides):
    """只改指定邊，其餘保留。sides 例：top='thin', bottom='double'。"""
    b = ws.cell(row=r, column=c).border
    kw = dict(top=b.top, bottom=b.bottom, left=b.left, right=b.right)
    for k, v in sides.items():
        kw[k] = Side(style=v)
    ws.cell(row=r, column=c).border = Border(**kw)


def _box(ws, r1, c1, r2, c2, style="medium"):
    """畫一個外框方框（含合併格：四邊都設在構成格上，Excel 才連續）。"""
    for c in range(c1, c2 + 1):
        _set_side(ws, r1, c, top=style)
        _set_side(ws, r2, c, bottom=style)
    for r in range(r1, r2 + 1):
        _set_side(ws, r, c1, left=style)
        _set_side(ws, r, c2, right=style)


def _table_bottom(ws, r, c1, c2, style="medium"):
    for c in range(c1, c2 + 1):
        _set_side(ws, r, c, bottom=style)

import simple_config as sc
import agency_cue as ac
from simple_style_master import StyleMaster
from simple_excel import thicken_hairlines, seal_grid

_WD = "一二三四五六日"
_WD_EN = ["M", "T", "W", "T", "F", "S", "S"]
_MON_EN = ["", "JAN", "FEB", "MAR", "APR", "MAY", "JUN",
           "JUL", "AUG", "SEP", "OCT", "NOV", "DEC"]
_LOGO_2008 = os.path.join(os.path.dirname(os.path.abspath(__file__)), "assets", "logo_2008.png")
_FILL_2008 = "FFFFFF99"
_FILL_CARAT = "FFFFFF00"


def _w(ws, r, c, v):
    ws.cell(row=r, column=c).value = v


def _fill(ws, r, c, color):
    ws.cell(row=r, column=c).fill = PatternFill("solid", fgColor=color) if color else PatternFill()


def _newsheet(wb, title, first):
    ws = wb.active if (first and wb.active.title == "Sheet") else wb.create_sheet(title)
    if first and ws.title == "Sheet":
        ws.title = title
    return ws


# =============================================================================
# 版型 B：2008 傳媒
# =============================================================================
def render_2008(wb, model, formulas):
    first = True
    is_wjf = model.combo_key.endswith("wjf")
    for sheet in model.sheets:
        ws = _newsheet(wb, sheet.title, first)
        first = False
        sm = StyleMaster(model.combo_key)
        nd = len(sheet.days)
        info = sm.apply(ws, nd)
        FIRST, LAST = info["first"], info["last"]
        # 全家母版「定價」欄(D)過窄，7 位數(如 1,440,000)在 22pt 會 ###；加寬(§7.4)
        if not is_wjf:
            ws.column_dimensions["D"].width = 33.0
        _2008_dateheader(ws, sheet, FIRST)
        if is_wjf:
            _fill_2008_wjf(ws, model, sheet, FIRST, LAST, formulas)
        else:
            _fill_2008_fam(ws, model, sheet, FIRST, LAST, formulas)
        _add_logo_2008(ws, LAST)
        thicken_hairlines(ws)
        seal_grid(ws, 8, sm.data_last, 1, LAST)    # 只密封資料表；下方另畫
        _table_bottom(ws, sm.data_last, 1, LAST, "double")   # 合計列底邊 double(全寬)
        for r in range(17, 21):                    # 費用四列 F:G 上下 thin
            for c in (6, 7):
                _set_side(ws, r, c, top="thin", bottom="thin")


def _2008_dateheader(ws, sheet, FIRST):
    for i, d in enumerate(sheet.days):
        c = FIRST + i
        if d.month_label:
            _w(ws, 8, c, _MON_EN[d.date.month])
        _w(ws, 9, c, d.date.day)
        _w(ws, 10, c, _WD[d.date.weekday()])
        _fill(ws, 10, c, _FILL_2008 if d.weekend else None)


def _abstract_2008(ws, model):
    _w(ws, 1, 1, "Client"); _w(ws, 1, 3, model.client)
    _w(ws, 2, 1, "Media Agency"); _w(ws, 2, 3, sc.COMPANY_2008)
    _w(ws, 3, 1, "Advertising Agency")
    _w(ws, 4, 1, "Product"); _w(ws, 4, 3, model.product)
    _w(ws, 5, 1, "Campaign"); _w(ws, 5, 3, model.campaign)
    _w(ws, 6, 1, "Period")
    _w(ws, 6, 3, f"{model.start:%Y.%m.%d}-{model.end:%Y.%m.%d}")
    for c, t in enumerate(["媒體型態", "地區", "播出時段", "定價", "單位", "次數", "合計",
                           "素材\n提供時間"], start=1):
        _w(ws, 8, c, t)


def _fees_2008(ws, sheet, formulas):
    f = sheet.fees
    labels = ["Budget (net)：", "AC %：", "5% Tax ：", "TOTAL ："]
    vals_f = ["=G11", "=G17*0.03", "=SUM(G17:G18)*0.05", "=SUM(G17:G19)"]
    vals_v = [sheet.budget, f["ac"], f["tax"], f["total"]]
    for i in range(4):
        _w(ws, 17 + i, 6, labels[i])
        _w(ws, 17 + i, 7, vals_f[i] if formulas else vals_v[i])


def _remarks_2008(ws, model, r0=21):
    for i, (txt, _red) in enumerate(model.remarks):
        _w(ws, r0 + i, 1, txt)


def _fill_2008_fam(ws, model, sheet, FIRST, LAST, formulas):
    _abstract_2008(ws, model)
    blk = sheet.blocks[0]
    main = blk.rows[0]
    bonus = [r for r in blk.rows if r.kind == "bonus"]
    FL, LL = get_column_letter(FIRST), get_column_letter(LAST)
    budget = sheet.budget

    _w(ws, 11, 1, ac.AGENCY_LABELS["2008傳媒"]["family"][0])
    _w(ws, 11, 2, "全省")
    _w(ws, 11, 3, "07:00-23:00")
    _w(ws, 11, 4, f"={main.list_per}*F11" if formulas else main.list_total)
    _w(ws, 11, 5, f"{main.seconds}秒")
    _w(ws, 11, 6, f"=SUM({FL}11:{LL}11)" if formulas else main.spots)
    _w(ws, 11, 7, budget)
    _w(ws, 11, 8, main.material)
    for i, v in enumerate(main.schedule):
        _w(ws, 11, FIRST + i, v)

    for bi, b in enumerate(bonus):
        r = 12 + bi
        _w(ws, r, 4, f"={b.list_per}*F{r}" if formulas else b.list_total)
        _w(ws, r, 6, b.spots)
        _w(ws, r, 7, sc.TXT_REBATE)
        _w(ws, r, FIRST, f"{b.spots}(走期內執行完畢)")

    # 合計列 14
    _w(ws, 14, 2, "合計")
    _w(ws, 14, 4, "=SUM(D11:D13)" if formulas else sum(
        [main.list_total] + [b.list_total for b in bonus]))
    _w(ws, 14, 6, "=SUM(F11:F13)" if formulas else main.spots + sum(b.spots for b in bonus))
    _w(ws, 14, 7, "=G11" if formulas else budget)
    for i in range(len(sheet.days)):
        col = FIRST + i
        cl = get_column_letter(col)
        _w(ws, 14, col, f"=SUM({cl}11:{cl}13)" if formulas else main.schedule[i])

    _fees_2008(ws, sheet, formulas)
    _remarks_2008(ws, model)


def _fill_2008_wjf(ws, model, sheet, FIRST, LAST, formulas):
    _abstract_2008(ws, model)
    rows = sheet.blocks[0].rows       # mag, sup, bonus, super_bonus
    mag, sup = rows[0], rows[1]
    bm = next(r for r in rows if r.kind == "bonus")
    bs = next(r for r in rows if r.kind == "super_bonus")
    FL, LL = get_column_letter(FIRST), get_column_letter(LAST)
    budget = sheet.budget

    _w(ws, 11, 1, ac.AGENCY_LABELS["2008傳媒"]["mag"][0])
    _w(ws, 12, 1, ac.AGENCY_LABELS["2008傳媒"]["super"][0])
    _w(ws, 13, 1, ac.AGENCY_LABELS["2008傳媒"]["mag"][0])
    _w(ws, 14, 1, ac.AGENCY_LABELS["2008傳媒"]["super"][0])
    _w(ws, 11, 2, "全省")            # B11:B12
    _w(ws, 13, 2, "全省")            # B13:B14
    _w(ws, 11, 3, "0900-2300")
    _w(ws, 12, 3, "0000-2400")
    _w(ws, 13, 3, "0900-2300")
    _w(ws, 14, 3, "0000-2400")
    _w(ws, 11, 4, f"={mag.list_per}*F11" if formulas else mag.list_total)
    _w(ws, 12, 4, sc.TXT_2008_ON_MAG)
    _w(ws, 13, 4, f"={bm.list_per}*F13" if formulas else bm.list_total)
    _w(ws, 14, 4, sc.TXT_2008_ON_MAG)
    _w(ws, 11, 5, f"{mag.seconds}秒")   # E11:E12
    _w(ws, 13, 5, f"{mag.seconds}秒")   # E13:E14
    _w(ws, 11, 6, f"=SUM({FL}11:{LL}11)" if formulas else mag.spots)
    _w(ws, 12, 6, f"=SUM({FL}12:{LL}12)" if formulas else sup.spots)
    _w(ws, 13, 6, bm.spots)
    _w(ws, 14, 6, bs.spots)
    _w(ws, 11, 7, budget)               # G11:G12
    _w(ws, 13, 7, sc.TXT_REBATE_SHORT)  # G13:G14
    _w(ws, 11, 8, mag.material)         # H11:H14
    for i in range(len(sheet.days)):
        _w(ws, 11, FIRST + i, mag.schedule[i])
        _w(ws, 12, FIRST + i, sup.schedule[i])
    _w(ws, 13, FIRST, bm.spots)         # 回饋列日期區合併只寫數字（§2.4）
    _w(ws, 14, FIRST, bs.spots)

    _w(ws, 15, 2, "合計")
    _w(ws, 15, 4, "=SUM(D11:D12)" if formulas else mag.list_total)
    _w(ws, 15, 6, "=SUM(F11:F14)" if formulas else mag.spots + sup.spots + bm.spots + bs.spots)
    _w(ws, 15, 7, "=G11" if formulas else budget)
    for i in range(len(sheet.days)):
        cl = get_column_letter(FIRST + i)
        _w(ws, 15, FIRST + i, f"=SUM({cl}11:{cl}12)" if formulas else mag.schedule[i] + sup.schedule[i])

    _fees_2008(ws, sheet, formulas)
    _remarks_2008(ws, model)


def _add_logo_2008(ws, out_last):
    if not os.path.exists(_LOGO_2008):
        return
    try:
        img = XLImage(_LOGO_2008)
        img.width, img.height = 150, 77
        anchor_col = max(out_last - 5, 1)
        ws.add_image(img, f"{get_column_letter(anchor_col)}1")
    except Exception:
        pass


# =============================================================================
# 版型 C：凱絡
# =============================================================================
def render_carat(wb, model, formulas):
    first = True
    is_wjf = model.combo_key.endswith("wjf")
    for sheet in model.sheets:
        ws = _newsheet(wb, sheet.title, first)
        first = False
        sm = StyleMaster(model.combo_key)
        nd = len(sheet.days)
        info = sm.apply(ws, nd)
        FIRST, LAST = info["first"], info["last"]
        # 加寬日期欄與「總價」欄，避免大數字顯示 ###（範本原欄寬過窄；§7.4）
        for c in range(FIRST, LAST + 1):
            ws.column_dimensions[get_column_letter(c)].width = 9.0
        ws.column_dimensions["I"].width = 17.0   # 總價欄(7~8 位數會計格式)
        # A 欄加寬到 24：Noto 字型下「全台【萬家福/樂家康】」在 21.5 會把「】」擠到第二行（§6）
        ws.column_dimensions["A"].width = 24.0
        _carat_dateheader(ws, sheet, FIRST)
        if is_wjf:
            _fill_carat_wjf(ws, model, sheet, FIRST, LAST, formulas)
        else:
            _fill_carat_fam(ws, model, sheet, FIRST, LAST, formulas)
        thicken_hairlines(ws)
        seal_grid(ws, 5, sm.data_last, 1, LAST)    # 只密封資料表(表頭5-7+資料)
        # 資料列以下（媒體總價值框/費用底線/簽核帶上緣/備註框/外框）一律由母版提供
        # （§4：母版下半部保留、不清空）；渲染器不再補畫，避免蓋成 medium 與範本不符。


def _carat_dateheader(ws, sheet, FIRST):
    for i, d in enumerate(sheet.days):
        c = FIRST + i
        if d.month_label:
            en = _MON_EN[d.date.month]
            _w(ws, 5, c, en[0] + en[1:].lower())    # Sep 式
        _w(ws, 6, c, d.date.day)
        _w(ws, 7, c, _WD_EN[d.date.weekday()])
        _fill(ws, 7, c, _FILL_CARAT if d.weekend else None)


def _carat_head(ws, model):
    _w(ws, 1, 1, "凱絡媒體服務(股)公司廣播媒體排期表")
    _w(ws, 4, 1, f"{model.start.year}年{model.start.month}月")
    _w(ws, 2, 9, "客   戶："); _w(ws, 2, 10, model.client)
    _w(ws, 3, 9, "產   品："); _w(ws, 3, 10, model.product)
    _w(ws, 4, 9, "日   期："); _w(ws, 4, 10, model.made_date.strftime("%Y/%m/%d"))
    for c, t in enumerate(["媒體別", "地區", "時段", "素材", "定價\n(檔/Net)", "市場價\n(檔/Net)",
                           "統一價\n(檔/Net)", "檔數", "總價", "專案價\n(Net)"], start=1):
        _w(ws, 5, c, t)


def _carat_datarow(ws, r, row, FIRST, LAST, formulas):
    FL, LL = get_column_letter(FIRST), get_column_letter(LAST)
    _w(ws, r, 4, f'{row.seconds}"CM')
    _w(ws, r, 5, row.list_per)
    _w(ws, r, 6, f"=E{r}*0.8" if formulas else row.market_per)
    _w(ws, r, 7, f"=E{r}*0.75" if formulas else row.uni_per)
    _w(ws, r, 8, f"=SUM({FL}{r}:{LL}{r})" if formulas else row.spots)
    _w(ws, r, 9, f"=G{r}*H{r}" if formulas else row.uni_total)
    if row.schedule:
        for i, v in enumerate(row.schedule):
            _w(ws, r, FIRST + i, v)


def _carat_signature(ws, out_last, r=21):
    _w(ws, r, 1, "部主管：_________ ")
    _w(ws, r, 4, "課主管：_________ ")
    _w(ws, r, 9, "媒體窗口：_________")
    o = 15 if 15 <= out_last else max(out_last - 3, 12)
    _w(ws, r, o, "承辦PM：_________")
    # O 落在日期區、樣式取自中間日期欄(無左對齊)，明確補左對齊對齊範本
    ws.cell(row=r, column=o).alignment = Alignment(horizontal="left", vertical="center")


def _carat_remarks(ws, model, r0=24):
    for i, (txt, _red) in enumerate(model.remarks):
        _w(ws, r0, 1, "備     註") if i == 0 else None
        _w(ws, r0 + i, 2, txt)


def _fill_carat_fam(ws, model, sheet, FIRST, LAST, formulas):
    _carat_head(ws, model)
    rows = sheet.blocks[0].rows       # main + 3 bonus
    budget = sheet.budget
    _w(ws, 8, 1, ac.AGENCY_LABELS["凱絡"]["family"][0])
    _w(ws, 8, 2, "全台\n全家便利商店")     # B8:B11
    _w(ws, 8, 3, "0700-2300")           # C8:C9
    _w(ws, 10, 3, "=C8" if formulas else "0700-2300")   # C10:C11
    for i, row in enumerate(rows):
        _carat_datarow(ws, 8 + i, row, FIRST, LAST, formulas)
    _w(ws, 8, 10, budget)               # 專案價 J8
    _w(ws, 9, 10, sc.TXT_REBATE)        # J9:J11 專案回饋
    # 媒體/優惠總價值 + 費用
    _w(ws, 14, 1, "媒體總價值(NET)")
    _w(ws, 14, 2, "=SUM(I8:I11)" if formulas else sheet.fees["media_value"])
    _w(ws, 15, 1, "優惠總價值(NET)")
    _w(ws, 15, 2, "=B14-J8" if formulas else sheet.fees["discount_value"])
    _w(ws, 14, 9, "Sub-Total"); _w(ws, 14, 10, "=J8" if formulas else budget)
    _w(ws, 15, 9, "A.C     3%"); _w(ws, 15, 10, "-")
    _w(ws, 16, 9, "VAT    5%")
    _w(ws, 16, 10, "=SUM(J14:J15)*5%" if formulas else sheet.fees["vat"])
    _w(ws, 17, 9, "Grand-Total")
    _w(ws, 17, 10, "=SUM(J14:J16)" if formulas else sheet.fees["grand"])
    _carat_signature(ws, LAST)
    _carat_remarks(ws, model)


def _fill_carat_wjf(ws, model, sheet, FIRST, LAST, formulas):
    _carat_head(ws, model)
    rows = sheet.blocks[0].rows       # mag, sup
    mag, sup = rows[0], rows[1]
    budget = sheet.budget
    _w(ws, 8, 1, ac.AGENCY_LABELS["凱絡"]["mag"][0])   # A8:A9
    _w(ws, 8, 2, "萬家福(量販)")
    _w(ws, 9, 2, "樂家康(超市)")
    _w(ws, 8, 3, "0900-2300 ")
    _w(ws, 9, 3, "0000-2400 ")
    _carat_datarow(ws, 8, mag, FIRST, LAST, formulas)
    _carat_datarow(ws, 9, sup, FIRST, LAST, formulas)
    _w(ws, 8, 10, budget)
    _w(ws, 9, 10, sc.TXT_CARAT_ON_MAG)
    # 費用（wjf 位置：I11..I14）
    _w(ws, 11, 9, "Sub-Total"); _w(ws, 11, 10, "=J8" if formulas else budget)
    _w(ws, 12, 9, "A.C     3%"); _w(ws, 12, 10, "-")
    _w(ws, 13, 9, "VAT    5%")
    _w(ws, 13, 10, "=SUM(J11:J12)*5%" if formulas else sheet.fees["vat"])
    _w(ws, 14, 9, "Grand-Total")
    _w(ws, 14, 10, "=SUM(J11:J13)" if formulas else sheet.fees["grand"])
    _w(ws, 14, 1, "媒體總價值(NET)")
    _w(ws, 14, 2, "=SUM(I8:I9)" if formulas else sheet.fees["media_value"])
    _w(ws, 15, 1, "優惠總價值(NET)")
    _w(ws, 15, 2, "=B14-J8" if formulas else sheet.fees["discount_value"])
    _carat_signature(ws, LAST)
    _carat_remarks(ws, model)
